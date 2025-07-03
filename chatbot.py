import streamlit as st
import openai
from response import openai_api_call
import nltk
from nltk.tokenize import word_tokenize
from nltk.corpus import stopwords
from fiber import FiberDBMS
from config import INDEX_FILE
import os
from docx import Document
from pptx import Presentation
import chardet
from PyPDF2 import PdfReader
import pandas as pd
from indexing import extract_keywords, detect_language

# NLTK data is now handled centrally in Arcanalte.py

def extract_content_from_file(uploaded_file):
    """Extracts text content from an uploaded file object."""
    file_extension = os.path.splitext(uploaded_file.name)[1].lower()
    content = ""
    try:
        if file_extension == ".txt":
            # Use chardet to detect encoding for robust text file reading
            raw_data = uploaded_file.getvalue()
            encoding = chardet.detect(raw_data)['encoding'] or 'utf-8'
            content = raw_data.decode(encoding, errors='replace')
        elif file_extension == ".docx":
            doc = Document(uploaded_file)
            content = "\n".join([para.text for para in doc.paragraphs])
        elif file_extension == ".pptx":
            presentation = Presentation(uploaded_file)
            all_texts = []
            for slide in presentation.slides:
                slide_texts = []
                if slide.shapes.title:
                    slide_texts.append(slide.shapes.title.text)
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        slide_texts.append(shape.text_frame.text)  # type: ignore
                all_texts.append("\n".join(slide_texts))
            content = "\n".join(all_texts)
        elif file_extension == ".pdf":
            reader = PdfReader(uploaded_file)
            content = "\n".join([page.extract_text() or '' for page in reader.pages])
        elif file_extension in [".xls", ".xlsx", ".csv"]:
            # For simplicity, we'll read CSV from the buffer; Excel might need saving to disk
            if "xls" in file_extension:
                # To handle xls/xlsx from stream, pandas may need the 'openpyxl' or 'xlrd' engine
                df = pd.read_excel(uploaded_file, engine='openpyxl' if file_extension == ".xlsx" else 'xlrd')
            else:
                df = pd.read_csv(uploaded_file)
            content = df.to_string()
        else:
            st.warning(f"Unsupported file format: {file_extension}. Cannot interpret this file.")
            return None
    except Exception as e:
        st.error(f"Failed to process and interpret file {uploaded_file.name}: {e}")
        return None
    return content

def chatbot_page():
    st.title("Chat With Arcana")

    # Ensure the database is loaded and available
    if 'dbms' not in st.session_state or not isinstance(st.session_state.dbms, FiberDBMS):
        st.warning("The database is not initialized. Please go to the 'Files' page and index your files first.")
        st.stop()

    dbms = st.session_state.dbms
    
    with st.sidebar:
        st.header("Upload a File")
        uploaded_file = st.file_uploader(
            "Upload a document to chat with it directly.",
            type=['txt', 'pdf', 'docx', 'pptx', 'csv', 'xlsx', 'xls']
        )
        if uploaded_file is not None:
            # Check if this file has been processed already to avoid reprocessing on every rerun
            if st.session_state.get('processed_file_name') != uploaded_file.name:
                with st.spinner(f"Interpreting {uploaded_file.name}..."):
                    file_content = extract_content_from_file(uploaded_file)
                    if file_content:
                        # Index the new file content into the database
                        with st.spinner(f"Indexing {uploaded_file.name}..."):
                            lines = file_content.split('\n')
                            for line in lines:
                                line = line.strip()
                                if line:
                                    lang = detect_language(line)
                                    keywords = extract_keywords(line, lang)
                                    dbms.add_entry(name=uploaded_file.name, content=line, tags=keywords)
                            dbms.save(INDEX_FILE) # Save the updated index
                            st.success(f"Successfully indexed {uploaded_file.name} into the database.")

                        # Add the file content as a system message for context, with priority instructions
                        context_message = (
                            f"The user has uploaded the file `{uploaded_file.name}`. "
                            f"For the user's next questions, you MUST prioritize the content of this file as the primary and ONLY source of information. "
                            f"Ignore any previous search results from the general document database. "
                            f"All answers must come directly from the file content provided below. "
                            f"Explicitly mention that you are answering based on the uploaded file."
                            f"\n\n--- FILE CONTENT ---\n{file_content}\n--- END FILE CONTENT ---"
                        )
                        st.session_state.messages.append({"role": "system", "content": context_message})
                        st.session_state.processed_file_name = uploaded_file.name
                        st.success(f"Successfully interpreted and prioritized {uploaded_file.name}! You can now ask questions about it.")
                    else:
                        st.session_state.processed_file_name = None # Reset if processing fails
        
        st.info("Clearing messages will remove the context from any uploaded file and revert to searching all documents.")

    if st.button("Clear All Messages"):
        # Reset messages but keep the initial system prompt
        init_messages()
        st.session_state.processed_file_name = None
        st.rerun()

    if "messages" not in st.session_state or not st.session_state.messages:
        init_messages()

    # Display existing conversation (excluding system messages)
    for message in st.session_state.messages:
        if message["role"] != "system":
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

    # User input area
    user_input = st.chat_input("Ask me anything about your documents...")

    # Dropdown for response type is maintained
    response_type = st.selectbox(
        "Choose a response type:",
        ["Normal", "IDX", "Math"],
        help="""
        - **Normal**: General conversation. The AI will use the search results as context but won't be strictly bound by them.
        - **IDX**: The AI will base its answer strictly on the content found in your indexed files.
        - **Math**: Specialized mode for mathematical queries.
        """
    )

    if user_input:
        st.session_state.messages.append({"role": "user", "content": user_input})
        with st.chat_message("user"):
            st.markdown(user_input)
        
        # If a file has NOT been processed, search the database for context.
        # If a file HAS been processed, its context is already in the messages, so we skip this.
        if st.session_state.get('processed_file_name') is None:
            with st.spinner("Searching for relevant information..."):
                stop_words = set(stopwords.words('english'))
                words = word_tokenize(user_input)
                keywords = [word for word in words if word.lower() not in stop_words and word.isalpha()]
                
                # Use the dbms instance from session state
                results = dbms.query(" ".join(keywords), top_n=min(20, max(1, len(keywords))))
                results = results[:5]

                assistant_reply = ""
                if results:
                    assistant_reply += "Here are the top results from your documents:\n\n"
                    for idx, result in enumerate(results, 1):
                        assistant_reply += f"**Result {idx} from `{result['name']}`:**\n"
                        assistant_reply += f"_{result['content']}_\n\n"
                else:
                    assistant_reply = "I couldn't find any specific information related to your query in the indexed documents."

                st.session_state.messages.append({"role": "system", "content": assistant_reply})

        with st.spinner("Arcana is thinking..."):
            try:
                # Make sure to initialize 'processed_file_name' if it doesn't exist
                if 'processed_file_name' not in st.session_state:
                    st.session_state.processed_file_name = None

                with st.chat_message("assistant"):
                    # Use st.write_stream to render the response in real-time
                    response_generator = openai_api_call(st.session_state.messages, response_type)
                    collected_chunks = []

                    def stream_and_collect():
                        for chunk in response_generator:
                            collected_chunks.append(chunk)
                            yield chunk

                    st.write_stream(stream_and_collect())
                    full_response = "".join(collected_chunks)
                
                # Append the full response to the message history
                st.session_state.messages.append({"role": "assistant", "content": full_response})
                st.rerun()

            except Exception as e:
                st.error(f"An error occurred while communicating with the AI: {e}")

def init_messages():
    """Initializes or resets the chat message history in the session state."""
    st.session_state.messages = [
        {"role": "assistant", "content": "Hey, I'm Arcana, your Indexademics AI assistant. Ask me anything about your indexed files!"},
        {"role": "system", "content": "You are a helpful AI assistant named Arcana. You will be provided with search results from a user's documents. Your task is to answer the user's questions based *only* on the provided text. Cite the source document's name for all information you provide, like this: `(Source: document_name.pdf)`. If the provided text does not contain the answer, state that clearly. Be friendly, cute, and helpful."},
    ]

"""Tests for the mixed markdown formatting parser."""

from docx import Document

from scripts.mixup import parse_markdown_to_word_runs


def test_parse_markdown_to_word_runs_handles_underline_formatting():
    """The parser should treat ``__text__`` as underlined content."""

    document = Document()
    paragraph = document.add_paragraph()

    parse_markdown_to_word_runs(paragraph, "Start __underlined__ end")

    underline_runs = [run for run in paragraph.runs if run.text == "underlined"]

    assert underline_runs, "Expected an underlined run for the wrapped text"
    assert all(run.font.underline for run in underline_runs), "Underline formatting was not applied"


def test_parse_markdown_to_word_runs_ignores_triple_underscores():
    """Triple underscores should not be misinterpreted as underline markers."""

    document = Document()
    paragraph = document.add_paragraph()

    parse_markdown_to_word_runs(paragraph, "Stay ___plain___ text")

    triple_run = next((run for run in paragraph.runs if "plain" in run.text), None)

    assert triple_run is not None, "Expected to retain the text wrapped in triple underscores"
    assert triple_run.font.underline is not True, "Triple underscores should not produce underlined runs"

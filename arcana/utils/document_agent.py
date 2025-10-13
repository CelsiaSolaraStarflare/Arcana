"""Utilities for Arcana's document analysis agent.

The document agent is responsible for reading full documents from the
indexed cache, producing structured summaries, and persisting the
summaries back into the Fiber database for future retrieval.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional

import chardet
import pandas as pd
from docx import Document
from pptx import Presentation
from PyPDF2 import PdfReader

from arcana.core.config import CACHE_DIR, INDEX_FILE
from arcana.utils.fiber import FiberDBMS
from arcana.utils.indexing import extract_keywords, detect_language
from arcana.utils.response import openai_api_call


@dataclass
class AgentSummaryResult:
    """Outcome of a document summarisation request."""

    file_name: str
    summary_text: str = ""
    summary_path: Optional[Path] = None
    created: bool = False
    reason: Optional[str] = None
    citation_path: Optional[str] = None

    @property
    def summary_name(self) -> Optional[str]:
        if self.summary_path is None:
            return None
        return self.summary_path.name


class DocumentAgent:
    """Agent that can read, reason about, and summarise documents."""

    def __init__(self, cache_dir: Path | str = CACHE_DIR, summary_subdir: str = "Summaries"):
        self.cache_dir = Path(cache_dir)
        self.summary_dir = self.cache_dir / summary_subdir
        self.summary_dir.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def summarise_document(self, file_name: str, dbms: FiberDBMS) -> AgentSummaryResult:
        """Ensure a summary exists for ``file_name`` and return it."""

        if not file_name:
            return AgentSummaryResult(file_name=file_name, reason="Missing file name")

        source_path = self._locate_file(file_name)
        if source_path is None:
            return AgentSummaryResult(
                file_name=file_name,
                reason="Original file could not be located in the cache directory.",
            )

        summary_path = self.summary_dir / f"{source_path.stem}_summary.md"

        if summary_path.exists():
            summary_text = summary_path.read_text(encoding="utf-8", errors="replace")
            result = AgentSummaryResult(
                file_name=file_name,
                summary_text=summary_text,
                summary_path=summary_path,
                created=False,
                citation_path=str(summary_path.relative_to(self.cache_dir)),
            )
        else:
            full_content = self._read_file_text(source_path)
            if not full_content.strip():
                return AgentSummaryResult(
                    file_name=file_name,
                    reason="File was empty or unreadable when attempting to summarise.",
                )

            summary_text = self._generate_summary(file_name, full_content)
            if not summary_text.strip():
                return AgentSummaryResult(
                    file_name=file_name,
                    reason="Language model returned an empty summary.",
                )

            summary_path.write_text(summary_text, encoding="utf-8")
            result = AgentSummaryResult(
                file_name=file_name,
                summary_text=summary_text,
                summary_path=summary_path,
                created=True,
                citation_path=str(summary_path.relative_to(self.cache_dir)),
            )

        self._ensure_summary_indexed(result, dbms)
        return result

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------
    def _locate_file(self, file_name: str) -> Optional[Path]:
        """Search ``CACHE_DIR`` for a file that matches ``file_name``.

        The lookup prefers an exact relative path match so that citations can
        reference nested directory structures. If that fails, the legacy
        basename search is used as a fallback for backwards compatibility.
        """

        if not file_name:
            return None

        candidate = Path(file_name)
        # Accept absolute paths that already point to a cached file.
        if candidate.is_absolute() and candidate.exists():
            return candidate

        # Attempt to resolve the file relative to the cache directory to
        # support nested paths stored in the search index.
        relative_candidate = self.cache_dir / candidate
        if relative_candidate.exists():
            return relative_candidate

        target_name = candidate.name
        for root, _, files in os.walk(self.cache_dir):
            if target_name in files:
                return Path(root) / target_name
        return None

    def _read_file_text(self, path: Path) -> str:
        extension = path.suffix.lower()
        try:
            if extension == ".txt":
                raw = path.read_bytes()
                encoding = chardet.detect(raw)["encoding"] or "utf-8"
                return raw.decode(encoding, errors="replace")
            if extension == ".docx":
                doc = Document(path)
                return "\n".join(para.text for para in doc.paragraphs)
            if extension == ".pptx":
                presentation = Presentation(path)
                slide_text: List[str] = []
                for slide in presentation.slides:
                    if slide.shapes.title:
                        slide_text.append(slide.shapes.title.text)
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            slide_text.append(shape.text_frame.text)  # type: ignore[attr-defined]
                return "\n".join(slide_text)
            if extension == ".pdf":
                reader = PdfReader(str(path))
                return "\n".join(page.extract_text() or "" for page in reader.pages)
            if extension in {".csv", ".xls", ".xlsx"}:
                if extension == ".csv":
                    df = pd.read_csv(path)
                else:
                    df = pd.read_excel(path)
                return df.to_string(index=False)
        except Exception as exc:  # pragma: no cover - defensive I/O guard
            return f"[Agent warning] Unable to read file {path.name}: {exc}"

        # Fall back to a safe read for unknown formats
        return path.read_text(encoding="utf-8", errors="replace")

    def _chunk_content(self, content: str, chunk_size: int = 3500) -> List[str]:
        cleaned = content.strip()
        if len(cleaned) <= chunk_size:
            return [cleaned]

        paragraphs = [para.strip() for para in re.split(r"\n{2,}", cleaned) if para.strip()]
        chunks: List[str] = []
        buffer = ""
        for para in paragraphs:
            candidate = f"{buffer}\n\n{para}".strip() if buffer else para
            if len(candidate) <= chunk_size:
                buffer = candidate
                continue

            if buffer:
                chunks.append(buffer)
                buffer = ""

            if len(para) <= chunk_size:
                buffer = para
                continue

            for start in range(0, len(para), chunk_size):
                chunks.append(para[start : start + chunk_size])

        if buffer:
            chunks.append(buffer)

        return chunks

    def _generate_summary(self, file_name: str, content: str) -> str:
        chunks = self._chunk_content(content)
        total = len(chunks)
        summaries: List[str] = []

        for idx, chunk in enumerate(chunks, start=1):
            messages: Iterable[dict] = [
                {
                    "role": "system",
                    "content": (
                        "You are Arcana's document analysis agent. Read the provided "
                        "document section in full and produce a concise yet thorough "
                        "summary. Capture the main ideas, important definitions, data "
                        "tables, and any actionable steps. Provide the summary as "
                        "clear bullet points."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Document: {file_name}\n"
                        f"Section {idx} of {total}. Summarise all important content.\n\n{chunk}"
                    ),
                },
            ]

            stream = openai_api_call(messages, "Long Text")
            chunk_summary = "".join(piece for piece in stream).strip()
            if chunk_summary:
                header = f"### Section {idx} of {total}\n{chunk_summary}"
                summaries.append(header)

        if not summaries:
            return ""

        combined = (
            f"# Agent Summary for {file_name}\n\n"
            "This summary was generated by Arcana's agent after reading the full document.\n\n"
            + "\n\n".join(summaries)
        )
        return combined.strip()

    def _ensure_summary_indexed(self, result: AgentSummaryResult, dbms: FiberDBMS) -> None:
        if not result.summary_text:
            return

        summary_name = result.summary_name or f"{result.file_name}_summary"

        existing = any(entry.get("name") == summary_name for entry in dbms.database)
        if existing:
            return

        lang = detect_language(result.summary_text)
        keywords = extract_keywords(result.summary_text, lang)
        dbms.add_entry(name=summary_name, content=result.summary_text, tags=keywords)
        dbms.save(INDEX_FILE)


def get_document_agent() -> DocumentAgent:
    """Return a singleton ``DocumentAgent`` instance."""

    if not hasattr(get_document_agent, "_instance"):
        get_document_agent._instance = DocumentAgent()
    return get_document_agent._instance  # type: ignore[attr-defined]


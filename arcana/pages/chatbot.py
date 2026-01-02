import base64
import mimetypes
from io import BytesIO

import openai
import streamlit as st
from arcana.utils.response import openai_api_call

# Ensure NLTK data is available before importing NLTK functions
import arcana.utils.nltk_setup
from arcana.utils.fiber import FiberDBMS
from arcana.core.config import CACHE_DIR
import os
import json
import datetime
import re
from datetime import date
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple
from docx import Document
from pptx import Presentation
import chardet
from PyPDF2 import PdfReader
import pandas as pd
from arcana.utils.indexing import extract_keywords, detect_language
from arcana.utils.document_agent import AgentSummaryResult, get_document_agent
from arcana.utils.web_search import search_brave_images, search_web
from arcana.utils import storage
from arcana.utils.auth import decrypt_json_from_file, encrypt_json_to_file


BASE_SYSTEM_PROMPT = (
    "You are a helpful AI assistant named Arcana. You will be provided with search results "
    "from a user's documents and, when enabled, web supplements. Prioritize the provided "
    "document snippets for answers and cite the source document's full relative path using "
    "MLA-style citations, for example: “Guide.” subject/guide.pdf. PDF file. If the provided "
    "text does not contain the answer but the question involves widely known general "
    "knowledge, answer accurately using your own knowledge. When web supplement results are "
    "provided, you may incorporate them but still cite the full URL using MLA conventions "
    "with an access date. Always conclude your reply with a line that begins with `Sources:` "
    "followed by a comma-separated list of the sources you used. If you had to rely solely on "
    "general knowledge and no citations are available, end with `Sources: No sources cited.` "
    "Be friendly, cute, and helpful."
)


def _unique_preserve_order(items: Iterable[str]) -> List[str]:
    """Return a list of unique items while preserving their original order."""

    seen = set()
    unique_items: List[str] = []
    for item in items:
        if item not in seen:
            unique_items.append(item)
            seen.add(item)
    return unique_items


def _parse_keyword_response(response_text: str) -> List[str]:
    """Parse the raw keyword response text returned by the language model."""

    if not response_text:
        return []

    cleaned = response_text.replace("Keywords:", "")
    # Split on commas, semicolons, or newlines and remove numbering/bullets
    raw_candidates = re.split(r"[\n,;]", cleaned)
    parsed: List[str] = []
    for candidate in raw_candidates:
        candidate = re.sub(r"^[\s\-•*\d\.]+", "", candidate).strip()
        if candidate:
            parsed.append(candidate)
    return _unique_preserve_order(parsed)


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def _encode_image_to_data_url(
    file_name: str,
    raw_bytes: bytes,
    mime_type: Optional[str] = None,
) -> str:
    """Return a base64 data URL for the provided image bytes."""

    inferred_mime = mime_type or mimetypes.guess_type(file_name)[0] or "image/png"
    b64_data = base64.b64encode(raw_bytes).decode("utf-8")
    return f"data:{inferred_mime};base64,{b64_data}"


def _data_url_to_bytes(data_url: str) -> Optional[bytes]:
    """Decode a base64 data URL into raw bytes."""

    if not data_url.startswith("data:"):
        return None

    try:
        _, encoded = data_url.split(",", 1)
        return base64.b64decode(encoded)
    except Exception:
        return None


def _render_message_content(content) -> None:
    """Render chat message content that may include text and images."""

    if isinstance(content, str):
        st.markdown(content)
        return

    if isinstance(content, list):
        for part in content:
            if isinstance(part, dict):
                part_type = part.get("type")
                if part_type == "text":
                    st.markdown(part.get("text", ""))
                elif part_type in {"image_url", "input_image"}:
                    image_info = part.get("image_url", {})
                    if isinstance(image_info, dict):
                        image_url = image_info.get("url", "")
                        caption = image_info.get("caption") or image_info.get("description")
                        if image_url:
                            if image_url.startswith("data:"):
                                image_bytes = _data_url_to_bytes(image_url)
                                if image_bytes is not None:
                                    st.image(BytesIO(image_bytes), caption=caption, use_column_width=True)
                                else:
                                    st.image(image_url, caption=caption, use_column_width=True)
                            else:
                                st.image(image_url, caption=caption, use_column_width=True)
                else:
                    st.markdown(str(part))
            else:
                st.markdown(str(part))
        return

    st.markdown(str(content))


def _message_content_to_plain_text(content) -> str:
    """Extract textual content from a chat message for metadata tasks."""

    if content is None:
        return ""

    if isinstance(content, str):
        return content

    if isinstance(content, list):
        fragments: List[str] = []
        for part in content:
            if isinstance(part, dict):
                part_type = part.get("type")
                if part_type == "text":
                    fragments.append(part.get("text", ""))
                elif part_type in {"image_url", "input_image"}:
                    fragments.append("[image]")
            else:
                fragments.append(str(part))
        return " ".join(fragment for fragment in fragments if fragment).strip()

    return str(content)


def _is_welcome_message(message: dict) -> bool:
    if not isinstance(message, dict):
        return False
    if message.get("role") != "assistant":
        return False
    content = _message_content_to_plain_text(message.get("content"))
    lowered = content.lower()
    return "arcana" in lowered and "assistant" in lowered


def _get_welcome_message() -> str:
    cached = st.session_state.get("welcome_message")
    if cached:
        return cached
    options = [
        "Hi, I'm Arcana. Ask me anything about your indexed files.",
        "Welcome back. What would you like to explore in your documents today?",
        "Hello! I can search your indexed files. What should we look for?",
        "Hey there—Arcana here. Ask a question and I’ll dig into your files.",
        "Hi! Ready when you are. What do you want to find in your documents?",
        "Welcome! Drop a question and I’ll pull the most relevant snippets.",
        "Hi! I’m Arcana, your document assistant. How can I help today?",
        "Hello! Upload a file or ask about your indexed content.",
        "Hi! Let’s find what you need. What’s your question?",
        "Welcome to Arcana. What would you like to learn from your files?",
        "Hey! Ask a question and I’ll search your document library.",
        "Hi! I can summarize, cite, and search your files. What’s up?",
        "Hello! Tell me what you need, and I’ll check your indexed sources.",
        "Welcome back! What topic should we pull from your documents?",
        "Hi! I’m ready to help—what are we looking for today?",
        "Hey there. Ask me about your files and I’ll cite sources.",
        "Hello! Let’s start with a question about your indexed materials.",
        "Welcome! I can search your docs or use web supplements if enabled.",
        "Hi! What do you want to know from your uploaded files?",
        "Hey! Arcana’s ready—ask a question to begin.",
    ]
    import random

    message = random.choice(options)
    st.session_state["welcome_message"] = message
    return message


def _build_sources_default_line(doc_results: Iterable[dict], web_results: Iterable[dict]) -> str:
    """Build a fallback `Sources:` line based on available document and web snippets."""

    today = date.today()
    accessed_date = _format_access_date(today)

    def _format_local_source(result: dict) -> str:
        raw_name = result.get("name")
        if not raw_name:
            return ""

        name = str(raw_name).strip()
        if not name:
            return ""

        path_candidate = str(result.get("citation_path") or name).strip()
        if not path_candidate:
            path_candidate = name

        path_text = path_candidate.replace("\\", "/")
        # Prefer a cleaned path representation when the entry looks like a stored file path.
        if "/" in path_text or Path(path_text).suffix:
            path_display = path_text
        else:
            path_display = path_text

        clean_title = _derive_title_from_name(path_candidate)
        descriptor = _describe_format(name, path_candidate)

        return f"“{clean_title}.” {path_display}. {descriptor}."

    def _format_web_source(result: dict) -> str:
        link = (result.get("link") or "").strip()
        if not link:
            return ""

        title = (result.get("title") or "").strip()
        safe_title = re.sub(r"\s+", " ", title).strip()
        safe_title = safe_title.replace("[", "(").replace("]", ")")

        if safe_title:
            return f"“{safe_title}.” {link}. Accessed {accessed_date}."
        return f"{link}. Accessed {accessed_date}."

    doc_citations: List[str] = []
    for result in doc_results:
        formatted = _format_local_source(result)
        if formatted:
            doc_citations.append(formatted)

    web_citations: List[str] = []
    for result in web_results:
        formatted = _format_web_source(result)
        if formatted:
            web_citations.append(formatted)

    ordered_sources = _unique_preserve_order([*doc_citations, *web_citations])
    if not ordered_sources:
        return "Sources: No sources cited."

    return "Sources: " + ", ".join(ordered_sources)


def _format_access_date(day: date) -> str:
    month_name = day.strftime("%B")
    day_number = day.strftime("%d").lstrip("0") or day.strftime("%d")
    return f"{day_number} {month_name} {day.year}"


def _derive_title_from_name(raw_name: str) -> str:
    cleaned = re.sub(r"\s*\(agent summary\)\s*$", "", raw_name, flags=re.IGNORECASE)
    path = Path(cleaned)
    stem = path.stem or cleaned
    if not stem:
        return "Document"

    spaced = re.sub(r"[_\-]+", " ", stem).strip()
    if not spaced:
        spaced = stem.strip()

    title = spaced.title()
    return title or "Document"


def _describe_format(raw_name: str, path_hint: Optional[str] = None) -> str:
    candidate = path_hint or raw_name
    lowered = candidate.lower()
    if "agent summary" in lowered or "agent summary" in raw_name.lower():
        return "Agent-generated summary"

    suffix = Path(candidate).suffix.lower().lstrip(".")
    format_map = {
        "pdf": "PDF file",
        "docx": "Word document",
        "pptx": "PowerPoint presentation",
        "txt": "Text file",
        "csv": "CSV file",
        "xls": "Excel workbook",
        "xlsx": "Excel workbook",
        "md": "Markdown document",
    }

    return format_map.get(suffix, "Document")


def search_brave(query: str, max_results: int = 3) -> List[dict]:
    """Fetch supplemental search results from Brave's Search API for the query."""

    return search_web(query=query, max_results=max_results)


def _looks_like_image_request(query: str) -> bool:
    """Heuristically detect if the user is asking for images."""

    if not query:
        return False

    normalized = re.sub(r"\s+", " ", query).strip().lower()
    if not normalized:
        return False

    image_keywords = [
        "image",
        "images",
        "photo",
        "photos",
        "picture",
        "pictures",
        "logo",
        "logos",
        "diagram",
        "diagrams",
        "chart",
        "charts",
        "graph",
        "graphs",
        "illustration",
        "illustrations",
        "icon",
        "icons",
        "screenshot",
        "screenshots",
        "wallpaper",
        "wallpapers",
        "art",
        "artwork",
        "flag",
        "flags",
        "map",
        "maps",
        "show me",
    ]

    return any(keyword in normalized for keyword in image_keywords)


def _ensure_sources_line(response_text: str, default_line: str) -> str:
    """Guarantee the assistant response ends with a `Sources:` line."""

    if not response_text:
        return default_line

    trimmed_response = response_text.rstrip("\n")
    lines = trimmed_response.splitlines()

    # Find the last non-empty line to check whether it already contains sources
    last_nonempty_line = ""
    for line in reversed(lines):
        if line.strip():
            last_nonempty_line = line.strip()
            break

    if last_nonempty_line.lower().startswith("sources:"):
        return trimmed_response + ("\n" if response_text.endswith("\n") else "")

    lower_sources_line = default_line.strip().lower()
    if lower_sources_line.startswith("sources:"):
        appended_line = default_line
    else:
        appended_line = "Sources: No sources cited."

    suffix = "\n" if not response_text.endswith("\n") else ""
    return response_text + suffix + appended_line


def generate_keywords_with_gpt(query: str, language: str, max_keywords: int = 15) -> Tuple[List[str], str]:
    """Generate a list of search keywords using the language model."""

    if not query:
        return [], ""

    prompt = [
        {
            "role": "system",
            "content": (
                "You craft focused keyword lists for searching a document database. "
                "Return a concise list (comma separated) of topical keywords and phrases that "
                "will help retrieve passages relevant to the user's request. "
                "Prioritize terms that are likely to appear in the source documents and stay on topic. "
                "Do not add explanations."
            ),
        },
        {
            "role": "user",
            "content": (
                "Generate search keywords for the following query. "
                f"Primary language hint: {language}. "
                f"Query: {query}"
            ),
        },
    ]

    try:
        stream = openai_api_call(prompt, "Normal")
        chunks: List[str] = []
        for piece in stream:
            chunks.append(piece)
        response_text = "".join(chunks).strip()
        keywords = _parse_keyword_response(response_text)
        if not keywords:
            return [], response_text
        return keywords[:max_keywords], response_text
    except Exception:
        return [], ""


def query_dbms_with_keywords(dbms: FiberDBMS, keywords: List[str], max_results: int = 5) -> Tuple[List[dict], List[Tuple[str, int]]]:
    """Query the DBMS with combinations of the provided keywords."""

    if not keywords:
        return [], []

    aggregated_results: List[dict] = []
    search_attempts: List[Tuple[str, int]] = []
    seen_entries = set()

    def add_results(query_text: str, new_results: List[dict]):
        if new_results is None:
            count = 0
        else:
            count = 0
            for result in new_results:
                identifier = (result.get("name"), result.get("content"))
                if identifier in seen_entries:
                    continue
                seen_entries.add(identifier)
                aggregated_results.append(result)
                count += 1
                if len(aggregated_results) >= max_results:
                    break
        search_attempts.append((query_text, count))
        return len(aggregated_results) >= max_results

    # First, try a combined query using the most important keywords.
    combined_query = " ".join(keywords[: min(5, len(keywords))])
    if combined_query:
        add_results(combined_query, dbms.query(combined_query, top_n=max_results))

    if len(aggregated_results) < max_results:
        for keyword in keywords:
            if add_results(keyword, dbms.query(keyword, top_n=max_results)):
                break

    return aggregated_results[:max_results], search_attempts


def inspect_database_for_keywords(
    dbms: FiberDBMS,
    keywords: List[str],
    max_commands: int = 8,
) -> Tuple[List[Dict[str, object]], List[str]]:
    """Inspect the database for files whose names match the provided keywords.

    This helper is used by the discrete agent mode to build an explicit audit
    trail of the internal lookups performed prior to generating a reply.
    """

    if not keywords:
        return [], []

    try:
        entries = getattr(dbms, "database", [])
    except AttributeError:
        return [], []

    if not isinstance(entries, list) or not entries:
        return [], []

    cache_root = Path(CACHE_DIR)
    normalized_keywords = _unique_preserve_order(
        [keyword.strip() for keyword in keywords if keyword and keyword.strip()]
    )

    audit_log: List[Dict[str, object]] = []
    matched_files: List[str] = []

    for keyword in normalized_keywords[:max_commands]:
        command = f'SCAN name CONTAINS "{keyword}"'
        keyword_lower = keyword.lower()
        command_matches: List[Dict[str, object]] = []

        for entry in entries:
            if not isinstance(entry, dict):
                continue
            name = entry.get("name")
            if not isinstance(name, str) or not name:
                continue
            if keyword_lower in name.lower():
                cache_path = cache_root / name
                match_info = {
                    "file": name,
                    "exists": cache_path.exists(),
                }
                command_matches.append(match_info)
                matched_files.append(name)

        audit_log.append({"command": command, "matches": command_matches})

    return audit_log, _unique_preserve_order(matched_files)


def run_discrete_toolkit(
    dbms: FiberDBMS,
    keywords: List[str],
    base_results: List[dict],
    max_content_queries: int = 6,
    max_content_results: int = 3,
    max_summaries: int = 3,
) -> Dict[str, object]:
    """Execute structured search and summarisation helpers for discrete mode."""

    audit_log, matched_by_name = inspect_database_for_keywords(dbms, keywords)

    normalized_keywords = _unique_preserve_order(keywords)
    content_logs: List[Dict[str, object]] = []
    extra_results: List[dict] = []

    seen_entries = {
        (result.get("name"), result.get("content"))
        for result in base_results
        if isinstance(result, dict)
    }

    matched_files = list(matched_by_name)

    for keyword in normalized_keywords[:max_content_queries]:
        hits = dbms.query(keyword, top_n=max_content_results) or []
        hit_names = [hit.get("name") for hit in hits if hit.get("name")]
        if hit_names:
            matched_files.extend(hit_names)
        content_logs.append({"keyword": keyword, "hits": hit_names})
        for hit in hits:
            identifier = (hit.get("name"), hit.get("content"))
            if identifier in seen_entries:
                continue
            seen_entries.add(identifier)
            extra_results.append(hit)

    summary_results: List[AgentSummaryResult] = []
    summary_logs: List[Dict[str, object]] = []

    summary_candidates = _unique_preserve_order(matched_files)[:max_summaries]
    if summary_candidates:
        agent = get_document_agent()
        for file_name in summary_candidates:
            summary = agent.summarise_document(file_name, dbms)
            if summary.summary_text:
                summary_results.append(summary)
                if summary.summary_name:
                    matched_files.append(summary.summary_name)
            status = "failed"
            if summary.summary_text:
                status = "created" if summary.created else "cached"
            summary_logs.append(
                {
                    "file": file_name,
                    "status": status,
                    "reason": summary.reason,
                    "summary_path": summary.citation_path,
                }
            )

    return {
        "audit_log": audit_log,
        "content_logs": content_logs,
        "summary_logs": summary_logs,
        "summary_results": summary_results,
        "matched_files": _unique_preserve_order(matched_files),
        "extra_results": extra_results,
    }

# NLTK data is now handled centrally in Arcanalte.py

# Chat History Management Functions

def get_chat_histories_dir():
    """Get the directory where chat histories are stored."""
    root = storage.get_user_data_dir()
    chat_dir = root / "chat_histories"
    chat_dir.mkdir(parents=True, exist_ok=True)
    return str(chat_dir)

def save_chat_history(session_name=None):
    """Save the current chat session to a file."""
    if "messages" not in st.session_state or not st.session_state.messages:
        st.warning("No messages to save!")
        return None
    
    if session_name is None:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        session_name = f"chat_{timestamp}"
    
    # Filter out system messages for cleaner history
    user_messages = [msg for msg in st.session_state.messages if msg["role"] != "system"]
    user_only_messages = [msg for msg in user_messages if msg.get("role") == "user"]

    # Generate a short tagline once there's at least one user message
    tagline = generate_chat_tagline(user_only_messages) if user_only_messages else ""

    chat_data = {
        "session_name": session_name,
        "tagline": tagline,
        "timestamp": datetime.datetime.now().isoformat(),
        "messages": user_messages,
        "processed_file_name": st.session_state.get('processed_file_name', None)
    }
    
    chat_dir = get_chat_histories_dir()
    key = storage.get_auth_key()
    suffix = ".enc" if key else ".json"
    file_path = os.path.join(chat_dir, f"{session_name}{suffix}")

    try:
        if key:
            encrypt_json_to_file(Path(file_path), chat_data, key)
        else:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(chat_data, f, indent=2, ensure_ascii=False)
        return file_path
    except Exception as e:
        st.error(f"Failed to save chat history: {e}")
        return None

def load_chat_history(file_path):
    """Load a chat session from a file."""
    try:
        key = storage.get_auth_key()
        if key and file_path.endswith(".enc"):
            chat_data = decrypt_json_from_file(Path(file_path), key)
        else:
            with open(file_path, 'r', encoding='utf-8') as f:
                chat_data = json.load(f)
        
        # Clear current messages completely and load from history
        st.session_state.messages = []
        
        # Add the initial system message first
        st.session_state.messages.append({
            "role": "system",
            "content": BASE_SYSTEM_PROMPT,
        })
        
        history_messages = chat_data.get("messages", [])
        if history_messages:
            filtered = []
            welcome_seen = False
            for msg in history_messages:
                if _is_welcome_message(msg):
                    if welcome_seen:
                        continue
                    welcome_seen = True
                filtered.append(msg)
            history_messages = filtered

        # Add the welcome message only if there are no user messages in history
        user_messages_in_history = [msg for msg in history_messages if msg.get("role") == "user"]
        has_welcome_in_history = any(_is_welcome_message(msg) for msg in history_messages)
        if not user_messages_in_history and not has_welcome_in_history:
            st.session_state.messages.append({
                "role": "assistant",
                "content": _get_welcome_message(),
            })
        
        # Load the chat history messages
        if history_messages:
            st.session_state.messages.extend(history_messages)
        
        # Restore processed file context if it exists
        st.session_state.processed_file_name = chat_data.get('processed_file_name', None)
        
        # Set the current session name to the loaded chat's name for continuous saving
        st.session_state.current_session_name = chat_data.get('session_name', None)
        
        return True
    except Exception as e:
        st.error(f"Failed to load chat history: {e}")
        return False

def get_available_chat_histories():
    """Get a list of available chat history files."""
    chat_dir = get_chat_histories_dir()
    histories = []
    
    key = storage.get_auth_key()
    for filename in os.listdir(chat_dir):
        if filename.endswith('.enc') and not key:
            continue
        if filename.endswith('.json') and key:
            continue
        if filename.endswith('.json') or filename.endswith('.enc'):
            file_path = os.path.join(chat_dir, filename)
            try:
                file_mtime = os.path.getmtime(file_path)

                if filename.endswith(".enc") and key:
                    chat_data = decrypt_json_from_file(Path(file_path), key)
                else:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        chat_data = json.load(f)
                timestamp = chat_data.get('timestamp')
                if not timestamp:
                    timestamp = datetime.datetime.fromtimestamp(file_mtime).isoformat()
                    chat_data['timestamp'] = timestamp
                tagline = chat_data.get('tagline', '')
                user_count = len([msg for msg in chat_data.get('messages', []) if msg.get('role') == 'user'])
                if not tagline and user_count > 0:
                    tagline = generate_chat_tagline(chat_data.get('messages', []))
                    if tagline:
                        chat_data['tagline'] = tagline
                if chat_data.get('timestamp') != timestamp or chat_data.get('tagline') != tagline:
                    try:
                        if filename.endswith(".enc") and key:
                            encrypt_json_to_file(Path(file_path), chat_data, key)
                        else:
                            with open(file_path, 'w', encoding='utf-8') as f:
                                json.dump(chat_data, f, indent=2, ensure_ascii=False)
                    except Exception:
                        pass
                histories.append({
                    'filename': filename,
                    'filepath': file_path,
                    'session_name': chat_data.get('session_name', filename[:-5]),
                    'tagline': tagline,
                    'timestamp': timestamp or 'Unknown',
                    'message_count': len(chat_data.get('messages', [])),
                    'modified_time': file_mtime
                })
            except:
                continue
    
    def _parse_timestamp(value: str) -> datetime.datetime:
        try:
            return datetime.datetime.fromisoformat(value)
        except Exception:
            return datetime.datetime.fromtimestamp(0)

    # Sort by chat timestamp to avoid reordering on metadata writes.
    histories.sort(
        key=lambda x: _parse_timestamp(x.get('timestamp') or ''),
        reverse=True
    )
    return histories

def generate_chat_tagline(messages):
    """Generate a short tagline from the conversation."""
    try:
        content = " ".join(
            _message_content_to_plain_text(msg.get("content"))
            for msg in messages
            if msg.get("role") == "user"
        )
        if not content:
            return ""
        prompt = [
            {
                "role": "system",
                "content": (
                    "You generate a concise chat title in 5 words or fewer. "
                    "Return only the title with no punctuation at the end."
                ),
            },
            {"role": "user", "content": content[:800]},
        ]
        stream = openai_api_call(prompt, "Normal")
        response_text = "".join(piece for piece in stream).strip()
        if response_text:
            words = response_text.split()
            return " ".join(words[:5])
        lang = detect_language(content)
        keywords = extract_keywords(content, lang)
        return " ".join(keywords[:3])
    except Exception:
        return ""

def auto_generate_chat_title(messages):
    """Generate a meaningful title for the chat based on the conversation content."""
    try:
        # Extract user messages for title generation
        user_messages = [
            _message_content_to_plain_text(msg.get('content'))
            for msg in messages
            if msg.get('role') == 'user'
        ]

        if not user_messages:
            return f"chat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"

        # Use first few user messages to generate title
        conversation_sample = " ".join(user_messages[:3])  # First 3 user messages
        
        # Limit the sample to avoid token limits
        if len(conversation_sample) > 500:
            conversation_sample = conversation_sample[:500] + "..."
        
        title_prompt = [
            {"role": "system", "content": "You are a helpful assistant that generates short, descriptive titles for chat conversations. Generate a concise title (3-6 words) that captures the main topic of the conversation. Only return the title, nothing else."},
            {"role": "user", "content": f"Generate a short title for this conversation: {conversation_sample}"}
        ]
        
        # Import here to avoid circular imports
        from response import openai_api_call
        
        title_generator = openai_api_call(title_prompt, "Normal")
        generated_title = "".join(title_generator).strip()
        
        # Clean up the title
        generated_title = generated_title.replace('"', '').replace("'", "")
        if len(generated_title) > 50:
            generated_title = generated_title[:50].strip()
        
        # Fallback if generation fails or is empty
        if not generated_title or len(generated_title) < 3:
            return f"chat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        return generated_title
        
    except Exception as e:
        print(f"Failed to generate title: {e}")
        return f"chat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"

def auto_save_current_chat():
    """Automatically save the current chat with an AI-generated title."""
    if "messages" not in st.session_state or not st.session_state.messages:
        return None
    
    # Check if there are any meaningful messages (user/assistant, not just system)
    meaningful_messages = [msg for msg in st.session_state.messages if msg["role"] in ["user", "assistant"]]
    
    if len(meaningful_messages) < 2:  # Need at least 1 user message and 1 assistant response
        return None
    
    with st.spinner("🤖 Generating title and saving chat..."):
        # Generate AI title
        auto_title = auto_generate_chat_title(st.session_state.messages)
        
        # Save with the generated title
        file_path = save_chat_history(auto_title)
        
        if file_path:
            st.success(f"💾 Auto-saved chat as: '{auto_title}'")
            return file_path
        
    return None

def continuous_save_chat():
    """Continuously update the current chat session without showing notifications."""
    if "messages" not in st.session_state or not st.session_state.messages:
        return None
    
    # Check if there are any meaningful messages
    meaningful_messages = [msg for msg in st.session_state.messages if msg["role"] in ["user", "assistant"]]
    
    if len(meaningful_messages) < 1:  # Need at least some conversation
        return None
    
    # Check if we already have a current session name, if not generate one
    if not hasattr(st.session_state, 'current_session_name') or not st.session_state.current_session_name:
        if len(meaningful_messages) >= 2:  # Only generate title when we have full exchange
            st.session_state.current_session_name = auto_generate_chat_title(st.session_state.messages)
        else:
            st.session_state.current_session_name = f"chat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"
    
    # Save silently in the background
    file_path = save_chat_history(st.session_state.current_session_name)
    return file_path

def delete_chat_history(file_path):
    """Delete a chat history file."""
    try:
        os.remove(file_path)
        return True
    except Exception as e:
        st.error(f"Failed to delete chat history: {e}")
        return False

def extract_content_from_file(uploaded_file):
    """Extracts text content from an uploaded file object."""
    file_extension = os.path.splitext(uploaded_file.name)[1].lower()
    content = ""
    try:
        if file_extension == ".txt":
            # Use chardet to detect encoding for robust text file reading
            raw_data = uploaded_file.getvalue()
            encoding = chardet.detect(raw_data)['encoding'] or 'utf-8'
            content = raw_data.decode(encoding, errors='replace')
        elif file_extension == ".docx":
            doc = Document(uploaded_file)
            content = "\n".join([para.text for para in doc.paragraphs])
        elif file_extension == ".pptx":
            presentation = Presentation(uploaded_file)
            all_texts = []
            for slide in presentation.slides:
                slide_texts = []
                if slide.shapes.title:
                    slide_texts.append(slide.shapes.title.text)
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        slide_texts.append(shape.text_frame.text)  # type: ignore
                all_texts.append("\n".join(slide_texts))
            content = "\n".join(all_texts)
        elif file_extension == ".pdf":
            reader = PdfReader(uploaded_file)
            content = "\n".join([page.extract_text() or '' for page in reader.pages])
        elif file_extension in [".xls", ".xlsx", ".csv"]:
            # For simplicity, we'll read CSV from the buffer; Excel might need saving to disk
            if "xls" in file_extension:
                # To handle xls/xlsx from stream, pandas may need the 'openpyxl' or 'xlrd' engine
                df = pd.read_excel(uploaded_file, engine='openpyxl' if file_extension == ".xlsx" else 'xlrd')
            else:
                df = pd.read_csv(uploaded_file)
            content = df.to_string()
        else:
            st.warning(f"Unsupported file format: {file_extension}. Cannot interpret this file.")
            return None
    except Exception as e:
        st.error(f"Failed to process and interpret file {uploaded_file.name}: {e}")
        return None
    return content

def chatbot_page():
    st.title("Chat With Arcana")

    st.session_state.setdefault("pending_vision_inputs", [])
    st.session_state.setdefault("vision_last_image_data", None)

    # Add streamlined ChatGPT-like styling
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600&display=swap');

    html, body, [class*="css"] {
        font-family: "Space Grotesk", sans-serif;
    }

    .block-container {
        max-width: 880px;
        padding-top: 1.25rem;
        padding-bottom: 3rem;
    }

    .sidebar .block-container {
        padding-top: 1rem;
        padding-bottom: 1rem;
    }

    [data-testid="stChatMessage"] {
        border-radius: 14px;
        padding: 12px 16px;
        margin: 10px 0;
        border: 1px solid rgba(15, 23, 42, 0.08);
        background: #f8fafc;
    }

    [data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] {
        font-size: 1rem;
        line-height: 1.6;
        color: #0f172a;
    }

    [data-testid="stChatMessage"][data-testid="stChatMessageUser"] {
        background: #e9f1ff;
        border-color: rgba(37, 99, 235, 0.2);
    }

    [data-testid="stChatInput"] textarea {
        border-radius: 14px !important;
        border: 1px solid rgba(15, 23, 42, 0.15) !important;
        padding: 12px 14px !important;
        font-size: 1rem !important;
        background: #ffffff !important;
    }

    .chat-history-item {
        background-color: #f7f7f8;
        border-radius: 8px;
        padding: 8px 12px;
        margin: 4px 0;
        border: 1px solid #e5e7eb;
    }

    .chat-history-item:hover {
        background-color: #eef2f7;
    }

    .new-chat-btn {
        background: linear-gradient(90deg, #0f172a 0%, #1e293b 100%);
        color: white;
        border: none;
        border-radius: 10px;
        padding: 12px;
        font-weight: 600;
        margin-bottom: 16px;
    }
    </style>
    """, unsafe_allow_html=True)

    # Initialize or load the database automatically
    if 'dbms' not in st.session_state or not isinstance(st.session_state.dbms, FiberDBMS):
        dbms = FiberDBMS()
        if storage.index_file_exists():
            with st.spinner("Loading existing database..."):
                try:
                    storage.load_dbms(dbms)
                    st.success("Database loaded successfully!")
                except Exception as e:
                    st.warning(f"Failed to load existing database: {e}. Starting with empty database.")
        else:
            st.info("No indexed files found. You can upload files directly or go to the 'Files' page to index documents.")
        st.session_state.dbms = dbms

    dbms = st.session_state.dbms
    
    with st.sidebar:
        # New Chat Button (prominent like ChatGPT)
        if st.button("➕ New Chat", key="new_chat_btn", use_container_width=True, type="primary"):
            # Auto-save current chat before clearing
            auto_save_current_chat()
            
            # Reset messages but keep the initial system prompt
            init_messages()
            st.session_state.processed_file_name = None
            # Reset the current session name for a fresh start
            st.session_state.current_session_name = None
            st.rerun()
        
        st.markdown("---")
        
        # Chat History Section
        histories = get_available_chat_histories()
        if histories:
            st.markdown("### Chats")
            for i, history in enumerate(histories[:5]):
                with st.container():
                    col1, col2 = st.columns([5, 1])
                    with col1:
                        display_title = history.get('tagline') or history['session_name']
                        if st.button(
                            display_title,
                            key=f"chat_{i}",
                            help=f"{history['message_count']} messages • {history['timestamp'][:10]}",
                            use_container_width=True,
                        ):
                            continuous_save_chat()
                            if load_chat_history(history['filepath']):
                                st.success("Chat loaded!")
                                st.rerun()
                        meta = f"{history['message_count']} messages • {history['timestamp'][:10]}"
                        if history.get('tagline'):
                            st.caption(f"{history['session_name']} · {meta}")
                        else:
                            st.caption(meta)
                    with col2:
                        if st.button("🗑️", key=f"del_{i}", help="Delete chat"):
                            if delete_chat_history(history['filepath']):
                                st.success("Deleted!")
                                st.rerun()

            if len(histories) > 5:
                with st.expander("All Chats", expanded=False):
                    for i, history in enumerate(histories[5:], start=5):
                        col1, col2 = st.columns([5, 1])
                        with col1:
                            display_title = history.get('tagline') or history['session_name']
                            if st.button(
                                display_title,
                                key=f"chat_all_{i}",
                                help=f"{history['message_count']} messages • {history['timestamp'][:10]}",
                                use_container_width=True,
                            ):
                                continuous_save_chat()
                                if load_chat_history(history['filepath']):
                                    st.success("Chat loaded!")
                                    st.rerun()
                            meta = f"{history['message_count']} messages • {history['timestamp'][:10]}"
                            if history.get('tagline'):
                                st.caption(f"{history['session_name']} · {meta}")
                            else:
                                st.caption(meta)
                        with col2:
                            if st.button("🗑️", key=f"del_all_{i}", help="Delete chat"):
                                if delete_chat_history(history['filepath']):
                                    st.success("Deleted!")
                                    st.rerun()
        else:
            st.caption("No chat history yet")
        
        # Session Info (compact)
        if "messages" in st.session_state and st.session_state.messages:
            meaningful_messages = [msg for msg in st.session_state.messages if msg["role"] in ["user", "assistant"]]
            if len(meaningful_messages) >= 2:
                st.caption(f"💬 {len(meaningful_messages)} messages • Auto-saves when starting new chat")

    # Remove the old clear button since it's now in sidebar as "New Chat"

    if "messages" not in st.session_state or not st.session_state.messages:
        init_messages()

    # Display existing conversation (excluding system messages)
    for message in st.session_state.messages:
        role = message.get("role")
        if role != "system":
            with st.chat_message(role):
                _render_message_content(message.get("content"))

    with st.container():
        controls_cols = st.columns([2.2, 1, 1], gap="medium")
        with controls_cols[0]:
            uploaded_file = st.file_uploader(
                "Add a document or image",
                type=['txt', 'pdf', 'docx', 'pptx', 'csv', 'xlsx', 'xls', 'png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp'],
                help="Supported formats: PDF, Word, PowerPoint, Text, Excel, CSV, common image types",
                label_visibility="collapsed",
            )
        with controls_cols[1]:
            web_supplement_enabled = st.checkbox(
                "Web supplement",
                help=(
                    "When enabled, Arcana fetches public web search results from Brave to"
                    " supplement your indexed documents."
                ),
                key="web_supplement_enabled",
            )
        with controls_cols[2]:
            response_type = st.selectbox(
                "Mode",
                ["Normal", "IDX", "Math", "Reasoning", "Discrete"],
                help="""
                **Normal**: General conversation with search context
                **IDX**: Strictly based on indexed files
                **Math**: Specialized for mathematical queries
                **Reasoning**: Uses a deep reasoning model that thinks step-by-step before replying
                **Discrete**: Runs explicit database scans before responding and verifies sources with careful reasoning
                """,
                label_visibility="collapsed",
            )

        if st.session_state.get('processed_file_name'):
            context_cols = st.columns([3, 1], gap="small")
            with context_cols[0]:
                st.info(f"📄 Chatting with: **{st.session_state.processed_file_name}**")
            with context_cols[1]:
                if st.button("Clear file context", use_container_width=True):
                    st.session_state.processed_file_name = None
                    st.session_state.messages = [
                        msg
                        for msg in st.session_state.messages
                        if not (
                            msg.get("role") == "system"
                            and "uploaded the file" in msg.get("content", "")
                        )
                    ]
                    st.rerun()

        if uploaded_file is not None:
            file_extension = os.path.splitext(uploaded_file.name)[1].lower()
            if file_extension in IMAGE_EXTENSIONS:
                image_bytes = uploaded_file.getvalue()
                data_url = _encode_image_to_data_url(
                    uploaded_file.name,
                    image_bytes,
                    getattr(uploaded_file, "type", None),
                )
                pending_images = st.session_state.setdefault("pending_vision_inputs", [])
                last_data_url = st.session_state.get("vision_last_image_data")
                if not any(img.get("data_url") == data_url for img in pending_images):
                    pending_images.append({
                        "name": uploaded_file.name,
                        "data_url": data_url,
                    })
                    st.session_state["vision_last_image_data"] = data_url
                    uploads_dir = os.path.join(CACHE_DIR, "Uploads")
                    os.makedirs(uploads_dir, exist_ok=True)
                    try:
                        file_path = os.path.join(uploads_dir, uploaded_file.name)
                        with open(file_path, "wb") as f:
                            f.write(image_bytes)
                        st.caption(f"Saved: {uploaded_file.name}")
                    except Exception as e:
                        st.warning(f"Could not save image to {uploads_dir}: {e}")
                    if last_data_url == data_url:
                        st.info("🖼️ Image re-queued for the next response.")
                    else:
                        st.success("🖼️ Image queued for analysis.")
                else:
                    st.info("🖼️ This image is already queued for the next response.")
            elif st.session_state.get('processed_file_name') != uploaded_file.name:
                with st.spinner(f"🔍 Processing {uploaded_file.name}..."):
                    file_content = extract_content_from_file(uploaded_file)
                    if file_content:
                        uploads_dir = os.path.join(CACHE_DIR, "Uploads")
                        os.makedirs(uploads_dir, exist_ok=True)

                        try:
                            file_path = os.path.join(uploads_dir, uploaded_file.name)
                            with open(file_path, "wb") as f:
                                f.write(uploaded_file.getbuffer())

                            txt_filename = os.path.splitext(uploaded_file.name)[0] + "_extracted.txt"
                            txt_path = os.path.join(uploads_dir, txt_filename)
                            with open(txt_path, "w", encoding="utf-8") as f:
                                f.write(file_content)

                            st.caption(f"Saved: {uploaded_file.name} + extracted text")
                        except Exception as e:
                            st.warning(f"Could not save file to {uploads_dir}: {e}")

                        with st.spinner("📚 Indexing content..."):
                            lines = file_content.split('\n')
                            for line in lines:
                                line = line.strip()
                                if line:
                                    lang = detect_language(line)
                                    keywords = extract_keywords(line, lang)
                                    dbms.add_entry(name=uploaded_file.name, content=line, tags=keywords)
                            storage.save_dbms(dbms)

                        context_message = (
                            f"The user has uploaded the file `{uploaded_file.name}`. "
                            f"For the user's next questions, you MUST prioritize the content of this file as the primary and ONLY source of information. "
                            f"Ignore any previous search results from the general document database. "
                            f"All answers must come directly from the file content provided below. "
                            f"Explicitly mention that you are answering based on the uploaded file."
                            f"\n\n--- FILE CONTENT ---\n{file_content}\n--- END FILE CONTENT ---"
                        )
                        st.session_state.messages.append({"role": "system", "content": context_message})
                        st.session_state.processed_file_name = uploaded_file.name
                        st.success(f"✅ {uploaded_file.name} indexed and ready for questions!")
                        st.rerun()
                    else:
                        st.session_state.processed_file_name = None

    pending_images = st.session_state.get("pending_vision_inputs", [])
    if pending_images:
        with st.expander("🖼️ Images queued for next message", expanded=False):
            for idx, image in enumerate(pending_images, start=1):
                caption = image.get("name") or f"Image {idx}"
                data_url = image.get("data_url", "")
                image_bytes = _data_url_to_bytes(data_url)
                if image_bytes is not None:
                    st.image(BytesIO(image_bytes), caption=caption, use_column_width=True)
                elif data_url:
                    st.image(data_url, caption=caption, use_column_width=True)

    # User input area
    user_input = st.chat_input("Ask me anything about your documents...")

    if user_input:
        st.session_state.pending_sources_default = "Sources: No sources cited."
        pending_images_to_send = list(st.session_state.get("pending_vision_inputs", []))
        if pending_images_to_send:
            user_content: List[dict] = []
            for image in pending_images_to_send:
                data_url = image.get("data_url")
                if not data_url:
                    continue
                user_content.append(
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": data_url,
                            "detail": "auto",
                        },
                    }
                )
            user_content.append({"type": "text", "text": user_input})
            user_message = {"role": "user", "content": user_content}
        else:
            user_message = {"role": "user", "content": user_input}

        st.session_state.messages.append(user_message)
        with st.chat_message("user"):
            _render_message_content(user_message["content"])
        st.session_state["pending_vision_inputs"] = []

        # If a file has NOT been processed, search the database for context.
        # If a file HAS been processed, its context is already in the messages, so we skip this.
        results: List[dict] = []
        search_attempts: List[Tuple[str, int]] = []
        keywords: List[str] = []
        keyword_generation_raw = ""
        keyword_source = "n/a"
        agent_summary_results: List[AgentSummaryResult] = []
        web_results: List[dict] = []
        image_results: List[dict] = []

        if st.session_state.get('processed_file_name') is None:
            with st.spinner("Searching for relevant information..."):
                lang = detect_language(user_input)
                keyword_source = "nltk"
                keyword_generation_raw = ""
                keywords = []
                if response_type in {"Normal", "Discrete"}:
                    gpt_keywords, keyword_generation_raw = generate_keywords_with_gpt(user_input, lang)
                    if gpt_keywords:
                        keywords = gpt_keywords
                        keyword_source = "language_model"
                if not keywords:
                    keywords = extract_keywords(user_input, lang)
                    keyword_source = "nltk" if keywords else keyword_source

                results = []
                search_attempts = []
                discrete_scan_logs: List[Dict[str, object]] = []
                discrete_content_logs: List[Dict[str, object]] = []
                discrete_summary_logs: List[Dict[str, object]] = []
                if keywords:
                    if keyword_source == "language_model":
                        results, search_attempts = query_dbms_with_keywords(dbms, keywords)
                    else:
                        query_text = " ".join(keywords[:20])
                        raw_results = dbms.query(
                            query_text,
                            top_n=min(20, max(1, len(keywords)))
                        )
                        raw_results = raw_results or []
                        results = raw_results[:5]
                        search_attempts = [(query_text, len(results))]

                    if response_type == "Discrete":
                        toolkit_output = run_discrete_toolkit(dbms, keywords, results)
                        discrete_scan_logs = toolkit_output.get("audit_log", [])
                        discrete_content_logs = toolkit_output.get("content_logs", [])
                        discrete_summary_logs = toolkit_output.get("summary_logs", [])
                        # Matched files are implicitly covered by the audit and summary logs.
                        agent_summary_results.extend(toolkit_output.get("summary_results", []))

                        extra_results = toolkit_output.get("extra_results", [])
                        if extra_results:
                            deduped_results: List[dict] = []
                            seen_pairs = set()
                            for item in list(results) + list(extra_results):
                                if not isinstance(item, dict):
                                    continue
                                identifier = (item.get("name"), item.get("content"))
                                if identifier in seen_pairs:
                                    continue
                                seen_pairs.add(identifier)
                                deduped_results.append(item)
                            results = deduped_results

                        if any([discrete_scan_logs, discrete_content_logs, discrete_summary_logs]):
                            with st.expander("Discrete mode tool audit", expanded=False):
                                if discrete_scan_logs:
                                    st.markdown("### File name scans")
                                    for record in discrete_scan_logs:
                                        st.markdown(f"**{record['command']}**")
                                        matches = record.get("matches", [])
                                        if matches:
                                            for match in matches:
                                                file_name = match.get("file", "?")
                                                exists = match.get("exists", False)
                                                status = "✅" if exists else "⚠️"
                                                st.markdown(f"- {status} `{file_name}`")
                                        else:
                                            st.markdown("- No matching files")
                                if discrete_content_logs:
                                    st.markdown("### Content searches")
                                    for log_entry in discrete_content_logs:
                                        hits = log_entry.get("hits", [])
                                        keyword = log_entry.get("keyword", "?")
                                        if hits:
                                            st.markdown(f"- `{keyword}` → {', '.join(f'`{name}`' for name in hits)}")
                                        else:
                                            st.markdown(f"- `{keyword}` → no content matches")
                                if discrete_summary_logs:
                                    st.markdown("### Document summaries")
                                    for summary_log in discrete_summary_logs:
                                        file_name = summary_log.get("file", "?")
                                        status = summary_log.get("status", "failed")
                                        summary_path = summary_log.get("summary_path")
                                        reason = summary_log.get("reason")
                                        details = f"status={status}"
                                        if summary_path:
                                            details += f", stored at `{summary_path}`"
                                        if reason and status == "failed":
                                            details += f" (reason: {reason})"
                                        st.markdown(f"- `{file_name}` → {details}")

                if web_supplement_enabled:
                    web_results = search_brave(user_input)
                    if _looks_like_image_request(user_input):
                        image_results = search_brave_images(user_input)

                assistant_reply_lines = [
                    "INTERNAL SEARCH CONTEXT (not visible to the user):",
                    "Evaluate the following snippets and only cite information that is accurate and relevant.",
                    "If no snippets answer the query, explain that the indexed files lacked the information and rely on well-known general knowledge to respond accurately.",
                    f"Keyword generation method: {keyword_source}",
                ]

                assistant_reply_lines.append(
                    "Use MLA-style citations that include each document's full relative path (e.g., “Guide.” subject/guide.pdf. PDF file.)."
                )

                if keywords:
                    assistant_reply_lines.append(
                        "Keywords considered: " + ", ".join(keywords[:20])
                    )
                if search_attempts:
                    assistant_reply_lines.append("Search attempts:")
                    for query_text, count in search_attempts:
                        assistant_reply_lines.append(
                            f"- `{query_text}` → {count} new snippet(s)"
                        )
                if keyword_generation_raw and keyword_source == "language_model":
                    assistant_reply_lines.append(
                        "Raw keyword suggestion response: " + keyword_generation_raw
                    )

                if response_type == "Discrete":
                    assistant_reply_lines.append(
                        "Discrete toolkit operations executed before composing the reply:"
                    )
                    if discrete_scan_logs:
                        for record in discrete_scan_logs:
                            command_text = record.get("command", "scan")
                            matches = record.get("matches", [])
                            if matches:
                                formatted_matches = ", ".join(
                                    f"{item.get('file', '?')}"
                                    + (" ✅" if item.get("exists") else " ⚠️ missing")
                                    for item in matches
                                )
                                assistant_reply_lines.append(
                                    f"- {command_text} → {formatted_matches}"
                                )
                            else:
                                assistant_reply_lines.append(
                                    f"- {command_text} → no matching files"
                                )
                    if discrete_content_logs:
                        for log_entry in discrete_content_logs:
                            keyword = log_entry.get("keyword", "?")
                            hits = log_entry.get("hits", [])
                            if hits:
                                formatted_hits = ", ".join(f"`{name}`" for name in hits)
                                assistant_reply_lines.append(
                                    f"- SEARCH content for `{keyword}` → {formatted_hits}"
                                )
                            else:
                                assistant_reply_lines.append(
                                    f"- SEARCH content for `{keyword}` → no matches"
                                )
                    if discrete_summary_logs:
                        for summary_log in discrete_summary_logs:
                            file_name = summary_log.get("file", "?")
                            status = summary_log.get("status", "failed")
                            if status in {"created", "cached"}:
                                destination = summary_log.get("summary_path")
                                if destination:
                                    assistant_reply_lines.append(
                                        f"- SUMMARISE `{file_name}` → {status} summary stored at `{destination}`"
                                    )
                                else:
                                    assistant_reply_lines.append(
                                        f"- SUMMARISE `{file_name}` → {status} summary"
                                    )
                            else:
                                reason = summary_log.get("reason")
                                if reason:
                                    assistant_reply_lines.append(
                                        f"- SUMMARISE `{file_name}` → failed ({reason})"
                                    )
                                else:
                                    assistant_reply_lines.append(
                                        f"- SUMMARISE `{file_name}` → failed"
                                    )
                    if not any([discrete_scan_logs, discrete_content_logs, discrete_summary_logs]):
                        assistant_reply_lines.append(
                            "- No discrete toolkit operations were executed because no reliable keywords were detected."
                        )
                    assistant_reply_lines.append(
                        "After reviewing the snippets and tool log, think through the reasoning quietly before composing the final reply."
                    )

                if web_supplement_enabled:
                    assistant_reply_lines.append(
                        "When citing web supplements, follow MLA style with the full URL and an access date."
                    )
                    if web_results:
                        assistant_reply_lines.append(
                            "Web supplement results from Brave (cite using the provided URLs):"
                        )
                        for result in web_results:
                            assistant_reply_lines.append(
                                f"- {result['title']} ({result['link']}): {result['snippet']}"
                            )
                    else:
                        assistant_reply_lines.append(
                            "Brave web supplement returned no usable results. If you rely on general knowledge, end with 'Sources: No sources cited.'"
                        )
                    if image_results:
                        assistant_reply_lines.append(
                            "Web image results from Brave (use URLs when citing images):"
                        )
                        for result in image_results:
                            image_url = result.get("image") or result.get("thumbnail") or result.get("link")
                            assistant_reply_lines.append(
                                f"- {result.get('title', 'Image result')} ({image_url})"
                            )

                if agent_summary_results:
                    assistant_reply_lines.append(
                        "Arcana's document summarizer analysed the following files prior to drafting the reply. Cite the original document names when you reference these summaries."
                    )
                    for summary in agent_summary_results:
                        assistant_reply_lines.append(
                            f"- `{summary.file_name}` (summary stored as `{summary.summary_name}`)"
                        )

                assistant_reply = "\n".join(assistant_reply_lines) + "\n\n"

                if results:
                    assistant_reply += "Here are the top results from your documents:\n\n"
                    for idx, result in enumerate(results, 1):
                        assistant_reply += f"**Result {idx} from `{result['name']}`:**\n"
                        assistant_reply += f"_{result['content']}_\n\n"
                elif keywords:
                    assistant_reply += (
                        "No specific passages were retrieved. Let the user know that the indexed documents did not contain "
                        "information matching their request, then answer using web supplements if available or accurate "
                        "general knowledge."
                    )
                else:
                    assistant_reply += (
                        "Keyword extraction failed. Inform the user that the system could not understand the query well "
                        "enough to search the documents."
                    )

                if agent_summary_results:
                    assistant_reply += "\nAgent-generated document overviews:\n\n"
                    for summary in agent_summary_results:
                        assistant_reply += f"### Agent summary for `{summary.file_name}`\n{summary.summary_text}\n\n"

                assistant_reply += (
                    "\nAlways end the final response with a 'Sources:' line listing every citation or 'Sources: No sources cited.'"
                )

                combined_results = list(results)
                if agent_summary_results:
                    combined_results.extend(
                        {
                            "name": summary.summary_name or f"{summary.file_name} (agent summary)",
                            "content": summary.summary_text,
                            "citation_path": summary.citation_path,
                        }
                        for summary in agent_summary_results
                    )

                st.session_state.pending_sources_default = _build_sources_default_line(combined_results, web_results)
                st.session_state.messages.append({"role": "system", "content": assistant_reply})
        with st.spinner("Arcana is thinking..."):
            try:
                # Make sure to initialize 'processed_file_name' if it doesn't exist
                if 'processed_file_name' not in st.session_state:
                    st.session_state.processed_file_name = None

                messages_to_send = list(st.session_state.messages)
                if response_type == "Reasoning":
                    messages_to_send.append({
                        "role": "system",
                        "content": (
                            "You are in deep reasoning mode. Think through the problem step by step before answering. "
                            "Provide a thorough explanation that breaks complex ideas into clear stages so the user can "
                            "understand the concept deeply."
                        )
                    })
                elif response_type == "Discrete":
                    messages_to_send.append(
                        {
                            "role": "system",
                            "content": (
                                "You are in discrete retrieval mode. Use the system context to understand which documents were"
                                " inspected. Plan your reasoning silently before replying, ensure every claim is supported by"
                                " the retrieved snippets or agent summaries, and be explicit when information is missing."
                                " Always end with a 'Sources:' line."
                            ),
                        }
                    )

                with st.chat_message("assistant"):
                    # Use st.write_stream to render the response in real-time
                    response_generator = openai_api_call(messages_to_send, response_type)
                    collected_chunks = []

                    def stream_and_collect():
                        for chunk in response_generator:
                            collected_chunks.append(chunk)
                            yield chunk

                    st.write_stream(stream_and_collect())
                    full_response = "".join(collected_chunks)
                    default_sources_line = st.session_state.get(
                        "pending_sources_default",
                        "Sources: No sources cited.",
                    )
                    ensured_response = _ensure_sources_line(full_response, default_sources_line)
                    if ensured_response != full_response:
                        extra_text = ensured_response[len(full_response):]
                        if extra_text:
                            st.markdown(extra_text)
                        full_response = ensured_response
                    st.session_state.pending_sources_default = "Sources: No sources cited."

                    if response_type == "Reasoning" and response_generator.has_reasoning():
                        with st.expander("Show Arcana's reasoning", expanded=False):
                            reasoning_text = response_generator.reasoning.replace("\n", "  \n")
                            st.markdown(reasoning_text or "(Reasoning trace was empty.)")

                    if image_results:
                        columns = st.columns(3)
                        for idx, image in enumerate(image_results):
                            target = columns[idx % len(columns)]
                            img_url = image.get("thumbnail") or image.get("image") or image.get("link")
                            caption = image.get("title") or "Image result"
                            with target:
                                if img_url:
                                    st.image(img_url, caption=caption, use_column_width=True)
                                link = image.get("link")
                                if link:
                                    st.caption(link)

                # Append the full response (plus image results) to the message history
                if image_results:
                    combined_content: List[dict] = [{"type": "text", "text": full_response}]
                    for image in image_results:
                        img_url = image.get("thumbnail") or image.get("image") or image.get("link")
                        if img_url:
                            combined_content.append(
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": img_url,
                                        "caption": image.get("title") or "Image result",
                                    },
                                }
                            )
                    st.session_state.messages.append({"role": "assistant", "content": combined_content})
                else:
                    st.session_state.messages.append({"role": "assistant", "content": full_response})
                
                # Continuously save the chat in the background
                continuous_save_chat()
                
                st.rerun()

            except Exception as e:
                st.error(f"An error occurred while communicating with the AI: {e}")

def init_messages():
    """Initializes or resets the chat message history in the session state."""
    st.session_state.messages = [
        {"role": "assistant", "content": _get_welcome_message()},
        {"role": "system", "content": BASE_SYSTEM_PROMPT},
    ]
    st.session_state.pending_vision_inputs = []
    st.session_state.vision_last_image_data = None

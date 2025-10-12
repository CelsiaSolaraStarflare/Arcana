"""Unified Textual Refiner page combining rewrite and instruction editing tools."""

from __future__ import annotations

import datetime
import io
import re
import uuid
from textwrap import shorten
from typing import Any, Dict, List, Optional

import streamlit as st
from docx import Document
from openai.types.chat import ChatCompletionMessageParam

from arcana.pages.components.style_guides import render_style_guide_controls
from arcana.utils.diff import generate_inline_diff_html, generate_version_diff_html
from arcana.utils.response import openai_api_call
from arcana.utils.style_guides import get_style_guide_rules
from arcana.utils.text_metrics import TextMetrics, calculate_text_metrics


REWRITE_PRESETS: List[Dict[str, str]] = [
    {
        "key": "formal",
        "label": "Formalize grammar",
        "instruction": "Polish grammar, eliminate errors, and deliver the message with a professional tone.",
    },
    {
        "key": "shorten",
        "label": "Shorten",
        "instruction": "Condense the passage by removing redundancy while preserving the essential ideas.",
    },
    {
        "key": "expand",
        "label": "Expand",
        "instruction": "Elaborate on the core ideas with more detail and helpful transitions without changing the intent.",
    },
    {
        "key": "tone",
        "label": "Adjust tone",
        "instruction": "Rephrase the text with a friendly, approachable tone that still sounds confident and clear.",
    },
    {
        "key": "bulletize",
        "label": "Bulletize",
        "instruction": "Rewrite the content as a concise, well-structured bulleted list summarizing the main points.",
    },
]


def _ensure_textual_refiner_state() -> None:
    """Initialize session state keys used by the Textual Refiner."""

    if "textual_refiner_draft" not in st.session_state:
        st.session_state.textual_refiner_draft = "Paste or draft your text here for refinement."
    if "textual_refiner_candidates" not in st.session_state:
        st.session_state.textual_refiner_candidates: List[Dict[str, Any]] = []
    if "textual_refiner_history" not in st.session_state:
        st.session_state.textual_refiner_history: List[Dict[str, Any]] = []
    if "textual_refiner_snippet" not in st.session_state:
        st.session_state.textual_refiner_snippet = ""
    if "textual_refiner_merge_text" not in st.session_state:
        st.session_state.textual_refiner_merge_text = ""
    if "textual_refiner_merge_origin" not in st.session_state:
        st.session_state.textual_refiner_merge_origin = ""
    if "textual_refiner_merge_sources" not in st.session_state:
        st.session_state.textual_refiner_merge_sources: List[str] = []
    if "textual_refiner_focus_prompt" not in st.session_state:
        st.session_state.textual_refiner_focus_prompt = ""
    if "textual_refiner_feedback" not in st.session_state:
        st.session_state.textual_refiner_feedback = ""
    if "active_style_guides" not in st.session_state:
        st.session_state.active_style_guides: List[str] = []
    if "textual_refiner_quality_checks" not in st.session_state:
        st.session_state.textual_refiner_quality_checks = ""
    if "textual_refiner_last_metrics" not in st.session_state:
        st.session_state.textual_refiner_last_metrics: Optional[Dict[str, float]] = None
    if "textual_refiner_prev_metrics" not in st.session_state:
        st.session_state.textual_refiner_prev_metrics: Optional[Dict[str, float]] = None
    if "textual_refiner_last_text" not in st.session_state:
        st.session_state.textual_refiner_last_text = st.session_state.textual_refiner_draft
    if "textual_refiner_export_include_citations" not in st.session_state:
        st.session_state.textual_refiner_export_include_citations = True
    if "textual_refiner_versions" not in st.session_state:
        st.session_state.textual_refiner_versions: List[Dict[str, Any]] = []
    if "textual_refiner_pending_revert" not in st.session_state:
        st.session_state.textual_refiner_pending_revert: Optional[Dict[str, Any]] = None
    if "textual_refiner_instruction_output" not in st.session_state:
        st.session_state.textual_refiner_instruction_output = ""
    if "textual_refiner_instruction_diff" not in st.session_state:
        st.session_state.textual_refiner_instruction_diff = ""
    if "textual_refiner_instruction_label" not in st.session_state:
        st.session_state.textual_refiner_instruction_label = ""


def _get_preset(preset_key: str) -> Optional[Dict[str, str]]:
    for preset in REWRITE_PRESETS:
        if preset["key"] == preset_key:
            return preset
    return None


def _split_sentences(text: str) -> List[str]:
    stripped = text.strip()
    if not stripped:
        return []
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", stripped) if s.strip()]
    if not sentences:
        return [stripped]
    return sentences


def _record_version_snapshot(text: str, label: str) -> None:
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    st.session_state.textual_refiner_versions.append(
        {
            "text": text,
            "timestamp": timestamp,
            "label": label,
        }
    )


def _log_change(source_text: str, rewritten_text: str, label: str, focus: str = "") -> None:
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    source_metrics = calculate_text_metrics(source_text)
    rewritten_metrics = calculate_text_metrics(rewritten_text)
    entry = {
        "timestamp": timestamp,
        "label": label,
        "focus": focus,
        "source_text": source_text,
        "rewritten_text": rewritten_text,
        "source_metrics": source_metrics.to_dict(),
        "metrics": rewritten_metrics.to_dict(),
    }
    st.session_state.textual_refiner_history.insert(0, entry)
    st.session_state.textual_refiner_prev_metrics = source_metrics.to_dict()
    st.session_state.textual_refiner_last_metrics = rewritten_metrics.to_dict()
    st.session_state.textual_refiner_last_text = rewritten_text
    _record_version_snapshot(source_text, f"Snapshot before: {label}")


def _find_candidate(candidate_id: str) -> Optional[Dict[str, Any]]:
    for candidate in st.session_state.textual_refiner_candidates:
        if candidate["id"] == candidate_id:
            return candidate
    return None


def _build_messages(
    source_text: str, preset: Dict[str, str], focus_prompt: str
) -> List[ChatCompletionMessageParam]:
    focus_suffix = (
        f"\nAdditional guidance: {focus_prompt.strip()}"
        if focus_prompt and focus_prompt.strip()
        else ""
    )
    user_content = (
        "Rewrite the following text according to the goal. Provide only the rewritten passage without commentary."
        f"\n\nGoal: {preset['instruction']}{focus_suffix}\n\nText:\n{source_text}"
    )
    style_rules = get_style_guide_rules(st.session_state.get("active_style_guides", []))
    system_content = (
        "You are an expert writing assistant who rewrites content to match a requested style."
        " Respond only with the fully rewritten text."
    )
    if style_rules:
        system_content += "\nIncorporate these style guide rules when rewriting:\n" + style_rules
    return [
        {
            "role": "system",
            "content": system_content,
        },
        {"role": "user", "content": user_content},
    ]


def _build_check_messages(text: str) -> List[ChatCompletionMessageParam]:
    """Compose messages for the quality-check assistant."""

    style_rules = get_style_guide_rules(st.session_state.get("active_style_guides", []))
    system_prompt = (
        "You are an editorial quality analyst."
        " Review drafts for clarity, grammar, tone consistency, and alignment with any supplied style guides."
        " Respond with a markdown-formatted report using bullet lists and short sections."
    )
    if style_rules:
        system_prompt += "\nRespect the following style guide directives when assessing the text:\n" + style_rules

    user_prompt = (
        "Assess the following draft. Summarize key strengths, list the most important opportunities"
        " for improvement, and call out any high-severity issues (grammar, consistency, inclusivity)."
        " Provide actionable suggestions and reference style guide expectations when relevant."
        f"\n\nDraft:\n{text}"
    )

    return [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]


def _build_metric_entries(
    current: TextMetrics, previous: Optional[TextMetrics]
) -> List[Dict[str, Optional[str]]]:
    """Return formatted metric rows for display and export."""

    entries: List[Dict[str, Optional[str]]] = []

    def delta_str(
        current_value: float,
        previous_value: Optional[float],
        *,
        precision: int = 0,
        suffix: str = "",
    ) -> Optional[str]:
        if previous_value is None:
            return None
        difference = current_value - previous_value
        if precision == 0:
            formatted = f"{difference:+.0f}"
        else:
            formatted = f"{difference:+.{precision}f}"
        return formatted + suffix

    prev_word = previous.word_count if previous else None
    prev_avg = previous.avg_sentence_length if previous else None
    prev_flesch = previous.flesch_reading_ease if previous else None
    prev_passive = previous.passive_voice_ratio if previous else None

    entries.append(
        {
            "label": "Word count",
            "value": str(current.word_count),
            "delta": delta_str(current.word_count, prev_word),
        }
    )
    entries.append(
        {
            "label": "Average sentence length",
            "value": f"{current.avg_sentence_length:.2f} words",
            "delta": delta_str(
                current.avg_sentence_length,
                prev_avg,
                precision=2,
                suffix=" words",
            ),
        }
    )
    entries.append(
        {
            "label": "Flesch Reading Ease",
            "value": f"{current.flesch_reading_ease:.2f}",
            "delta": delta_str(current.flesch_reading_ease, prev_flesch, precision=2),
        }
    )
    passive_value = (
        f"{current.passive_voice_count} sentences ({current.passive_voice_ratio:.2f}%)"
    )
    passive_delta = delta_str(
        current.passive_voice_ratio, prev_passive, precision=2, suffix=" pts"
    )
    entries.append(
        {
            "label": "Passive voice",
            "value": passive_value,
            "delta": passive_delta,
        }
    )

    return entries


def _create_docx_report(
    *,
    text: str,
    metrics_entries: List[Dict[str, Optional[str]]],
    quality_check: str,
    style_guides: List[str],
    include_citations: bool,
    history: List[Dict[str, Any]],
) -> bytes:
    document = Document()
    document.add_heading("Arcana Textual Refiner Report", level=1)
    generated_ts = datetime.datetime.now().strftime("Generated on %Y-%m-%d %H:%M")
    document.add_paragraph(generated_ts)

    if style_guides:
        document.add_paragraph("Style guides applied: " + ", ".join(style_guides))

    document.add_heading("Current Draft", level=2)
    document.add_paragraph(text if text.strip() else "(Draft is currently empty.)")

    document.add_heading("Quality Metrics", level=2)
    table = document.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Metric"
    hdr_cells[1].text = "Current"
    hdr_cells[2].text = "Δ"
    for entry in metrics_entries:
        row_cells = table.add_row().cells
        row_cells[0].text = entry["label"]
        row_cells[1].text = entry["value"]
        row_cells[2].text = entry.get("delta") or ""

    document.add_heading("Quality Check", level=2)
    if quality_check and quality_check.strip():
        for raw_line in quality_check.strip().splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith("#"):
                heading_level = min(line.count("#"), 5)
                document.add_heading(line.lstrip("# "), level=min(heading_level + 1, 5))
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                document.add_paragraph(line)
    else:
        document.add_paragraph("Quality check has not been run yet.")

    document.add_heading("Applied Fixes", level=2)
    if history:
        for idx, entry in enumerate(history, start=1):
            prefix = f"[Fix {idx}] " if include_citations else ""
            label = entry.get("label", "Unnamed fix")
            timestamp = entry.get("timestamp", "")
            focus = entry.get("focus")
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(f"{prefix}{label}")
            if timestamp:
                paragraph.add_run(f" — {timestamp}")
            if focus:
                paragraph.add_run(f" (Focus: {focus})")
    else:
        document.add_paragraph("No fixes have been applied yet.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _create_markdown_report(
    *,
    text: str,
    metrics_entries: List[Dict[str, Optional[str]]],
    quality_check: str,
    style_guides: List[str],
    include_citations: bool,
    history: List[Dict[str, Any]],
) -> str:
    lines: List[str] = []
    lines.append("# Arcana Textual Refiner Report")
    lines.append("")
    lines.append(f"_Generated on {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}_")
    if style_guides:
        lines.append("")
        lines.append("**Style guides applied:** " + ", ".join(style_guides))

    lines.append("")
    lines.append("## Current Draft")
    lines.append("")
    lines.append(text.strip() or "(Draft is currently empty.)")

    lines.append("")
    lines.append("## Quality Metrics")
    lines.append("")
    lines.append("| Metric | Current | Δ |")
    lines.append("| --- | --- | --- |")
    for entry in metrics_entries:
        delta = entry.get("delta") or ""
        lines.append(f"| {entry['label']} | {entry['value']} | {delta} |")

    lines.append("")
    lines.append("## Quality Check")
    lines.append("")
    if quality_check and quality_check.strip():
        lines.append(quality_check.strip())
    else:
        lines.append("Quality check has not been run yet.")

    lines.append("")
    lines.append("## Applied Fixes")
    lines.append("")
    if history:
        for idx, entry in enumerate(history, start=1):
            prefix = f"[Fix {idx}] " if include_citations else ""
            label = entry.get("label", "Unnamed fix")
            timestamp = entry.get("timestamp", "")
            focus = entry.get("focus")
            detail = f" — {timestamp}" if timestamp else ""
            focus_suffix = f" _(Focus: {focus})_" if focus else ""
            lines.append(f"- {prefix}{label}{detail}{focus_suffix}")
    else:
        lines.append("- No fixes have been applied yet.")

    return "\n".join(lines)


def _generate_rewrites(selected_keys: List[str], focus_prompt: str) -> None:
    source_text = st.session_state.textual_refiner_draft
    st.session_state.textual_refiner_candidates = []
    st.session_state.textual_refiner_merge_origin = source_text
    st.session_state.textual_refiner_merge_sources = []

    for preset_key in selected_keys:
        preset = _get_preset(preset_key)
        if not preset:
            continue

        container = st.container()
        container.markdown(f"#### {preset['label']} (generating…)")
        stream_placeholder = container.empty()

        try:
            messages = _build_messages(source_text, preset, focus_prompt)
            response_generator = openai_api_call(messages, "Normal")  # type: ignore[arg-type]
            collected_response = ""
            for chunk in response_generator:
                collected_response += chunk
                stream_placeholder.markdown(
                    f"**{preset['label']}**\n\n{collected_response}"
                )

            candidate_text = collected_response.strip()
            if candidate_text:
                candidate = {
                    "id": str(uuid.uuid4()),
                    "key": preset_key,
                    "label": preset["label"],
                    "text": candidate_text,
                    "source": source_text,
                    "focus": focus_prompt.strip(),
                    "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
                st.session_state.textual_refiner_candidates.append(candidate)
                stream_placeholder.markdown(
                    f"**{preset['label']}**\n\n{candidate_text}"
                )
            else:
                stream_placeholder.warning(
                    "The model did not return any rewritten text."
                )
        except Exception as exc:  # pylint: disable=broad-except
            stream_placeholder.error(f"Unable to generate rewrite: {exc}")


def _add_sentences_to_merge(candidate_id: str, selection_key: str) -> None:
    candidate = _find_candidate(candidate_id)
    if not candidate:
        return
    selected_sentences = st.session_state.get(selection_key, [])
    if not selected_sentences:
        return

    additions = "\n".join(selected_sentences)
    existing = st.session_state.textual_refiner_merge_text.strip()
    st.session_state.textual_refiner_merge_text = (
        f"{existing}\n{additions}".strip() if existing else additions
    )

    if not st.session_state.textual_refiner_merge_origin:
        st.session_state.textual_refiner_merge_origin = candidate["source"]
    if (
        candidate["label"]
        not in st.session_state.textual_refiner_merge_sources
    ):
        st.session_state.textual_refiner_merge_sources.append(
            candidate["label"]
        )

    st.session_state[selection_key] = []


def _clear_merge_workspace() -> None:
    st.session_state.textual_refiner_merge_text = ""
    st.session_state.textual_refiner_merge_sources = []
    st.session_state.textual_refiner_merge_origin = (
        st.session_state.textual_refiner_draft
    )


def _finalize_merge() -> None:
    merged_text = st.session_state.textual_refiner_merge_text.strip()
    if not merged_text:
        return

    source_text = (
        st.session_state.textual_refiner_merge_origin
        or st.session_state.textual_refiner_draft
    )
    sources = ", ".join(st.session_state.textual_refiner_merge_sources)
    label = "Merged selection" if not sources else f"Merged selection ({sources})"

    _log_change(source_text, merged_text, label)
    st.session_state.textual_refiner_draft = merged_text
    st.session_state.textual_refiner_feedback = (
        "Merged rewrite applied to the draft."
    )
    st.session_state.textual_refiner_instruction_output = ""
    st.session_state.textual_refiner_instruction_diff = ""
    _clear_merge_workspace()
    st.rerun()


def _accept_candidate(candidate_id: str) -> None:
    candidate = _find_candidate(candidate_id)
    if not candidate:
        return

    _log_change(
        candidate["source"],
        candidate["text"],
        candidate["label"],
        candidate.get("focus", ""),
    )
    st.session_state.textual_refiner_draft = candidate["text"]
    st.session_state.textual_refiner_feedback = (
        f"Applied {candidate['label']} rewrite to the draft."
    )
    st.session_state.textual_refiner_instruction_output = ""
    st.session_state.textual_refiner_instruction_diff = ""
    _clear_merge_workspace()
    st.rerun()


def _run_instruction_edit(instruction: str) -> None:
    draft = st.session_state.textual_refiner_draft
    if not draft.strip() or not instruction.strip():
        return

    style_rules = get_style_guide_rules(st.session_state.get("active_style_guides", []))
    system_content = (
        "You are an expert editor. You will be given a piece of text and an instruction."
        " Your task is to rewrite the text based only on the instruction. Return nothing but"
        " the fully rewritten text."
    )
    if style_rules:
        system_content += "\nFollow these style guide rules while editing:\n" + style_rules

    messages: List[ChatCompletionMessageParam] = [
        {
            "role": "system",
            "content": system_content,
        },
        {
            "role": "user",
            "content": (
                f"Instruction: {instruction}\n\nText to edit:\n---\n{draft}"
            ),
        },
    ]

    try:
        response_generator = openai_api_call(messages, "Normal")  # type: ignore[arg-type]
        full_response = "".join(list(response_generator))
        cleaned_response = full_response.strip()
        if cleaned_response:
            st.session_state.textual_refiner_instruction_output = cleaned_response
            st.session_state.textual_refiner_instruction_label = instruction
            st.session_state.textual_refiner_instruction_diff = generate_inline_diff_html(
                draft, cleaned_response
            )
        else:
            st.warning(
                "The model returned an empty response. Try refining your instruction."
            )
    except Exception as exc:  # pylint: disable=broad-except
        st.error(f"An error occurred during editing: {exc}")


def _discard_instruction_edit() -> None:
    st.session_state.textual_refiner_instruction_output = ""
    st.session_state.textual_refiner_instruction_diff = ""
    st.session_state.textual_refiner_instruction_label = ""


def _apply_instruction_edit() -> None:
    proposed = st.session_state.textual_refiner_instruction_output.strip()
    if not proposed:
        return

    instruction = st.session_state.textual_refiner_instruction_label or "Instruction edit"
    label = f"Instruction edit — {shorten(instruction, width=48, placeholder='…')}"
    source = st.session_state.textual_refiner_draft

    _log_change(source, proposed, label, st.session_state.textual_refiner_focus_prompt)
    st.session_state.textual_refiner_draft = proposed
    st.session_state.textual_refiner_feedback = "Instruction-led edit applied to the draft."
    _discard_instruction_edit()
    st.rerun()


def _queue_revert(version_index: int) -> None:
    versions = st.session_state.textual_refiner_versions
    if version_index < 0 or version_index >= len(versions):
        return
    target = versions[version_index]
    st.session_state.textual_refiner_pending_revert = {
        "index": version_index,
        "text": target["text"],
        "label": target["label"],
        "timestamp": target["timestamp"],
    }
    st.rerun()


def _confirm_revert() -> None:
    pending = st.session_state.textual_refiner_pending_revert
    if not pending:
        return

    current_text = st.session_state.textual_refiner_draft
    target_text = pending["text"]
    label = pending.get("label", "previous version")

    _log_change(current_text, target_text, f"Reverted to {label}")
    st.session_state.textual_refiner_draft = target_text
    st.session_state.textual_refiner_feedback = (
        f"Reverted draft to {label}."
    )
    _discard_instruction_edit()
    st.session_state.textual_refiner_pending_revert = None
    st.rerun()


def _cancel_revert() -> None:
    st.session_state.textual_refiner_pending_revert = None
    st.rerun()


def textual_refiner_page() -> None:
    """Render the unified Textual Refiner interface."""

    _ensure_textual_refiner_state()

    st.title("🧠 Textual Refiner")
    st.write(
        "Generate rewrites, run targeted instruction edits, and manage revisions from a single workspace."
    )

    if st.session_state.textual_refiner_feedback:
        st.success(st.session_state.textual_refiner_feedback)
        st.session_state.textual_refiner_feedback = ""

    if not st.session_state.textual_refiner_merge_text:
        st.session_state.textual_refiner_merge_origin = (
            st.session_state.textual_refiner_draft
        )

    pending_revert = st.session_state.textual_refiner_pending_revert
    if pending_revert:
        st.warning("⚠️ **Confirm Revert Action**")
        st.write(
            f"Are you sure you want to revert to '{pending_revert.get('label', 'selected version')}'?"
        )
        st.info(
            "Your current draft will be saved as a snapshot before applying the revert so you can undo it later."
        )
        col_confirm, col_cancel, _ = st.columns([1, 1, 3])
        with col_confirm:
            st.button("✅ Yes, revert", type="primary", on_click=_confirm_revert)
        with col_cancel:
            st.button("❌ Cancel", on_click=_cancel_revert)
        st.markdown("---")

    diagnostics_enabled = st.session_state.get(
        "textual_refiner_enable_diagnostics",
        st.session_state.get("refiner_enable_diagnostics", True),
    )
    quality_checks_enabled = st.session_state.get(
        "textual_refiner_enable_quality_checks",
        st.session_state.get("refiner_enable_quality_checks", True),
    )

    tab_studio, tab_diagnostics, tab_history = st.tabs(
        ["Studio", "Diagnostics", "History"]
    )

    with tab_studio:
        col_main, col_side = st.columns([3, 2])

        with col_main:
            st.subheader("Working draft")
            st.session_state.textual_refiner_draft = st.text_area(
                "Draft text",
                value=st.session_state.textual_refiner_draft,
                height=360,
                placeholder="Enter the passage you would like Arcana to refine.",
            )

            st.session_state.textual_refiner_focus_prompt = st.text_input(
                "Optional context (tone, audience, key details)",
                value=st.session_state.textual_refiner_focus_prompt,
                placeholder="e.g., 'Aim for a confident but friendly voice.'",
            )

        with col_side:
            active_guides = render_style_guide_controls(context_key="textual_refiner")
            if active_guides:
                st.caption("Active guides: " + ", ".join(active_guides))

            st.subheader("Quick instruction edit")
            instruction = st.text_input(
                "Editing instruction",
                key="textual_refiner_instruction_input",
                placeholder="e.g., 'Make the introduction warmer without adding length.'",
            )
            instruction_disabled = (
                not st.session_state.textual_refiner_draft.strip()
                or not instruction.strip()
            )
            st.button(
                "✨ Run instruction",
                type="primary",
                disabled=instruction_disabled,
                on_click=_run_instruction_edit,
                args=(instruction,),
            )

            st.subheader("Rewrite presets")
            preset_options = [preset["key"] for preset in REWRITE_PRESETS]
            selected_keys = st.multiselect(
                "Select rewrite styles",
                options=preset_options,
                format_func=lambda key: _get_preset(key)["label"] if _get_preset(key) else key,
                help="Choose one or more presets to generate multiple candidates at once.",
            )

            run_disabled = (
                not st.session_state.textual_refiner_draft.strip() or not selected_keys
            )
            st.button(
                "🚀 Generate rewrites",
                type="primary",
                disabled=run_disabled,
                on_click=_generate_rewrites,
                args=(selected_keys, st.session_state.textual_refiner_focus_prompt),
            )

            st.subheader("Snippet clipboard")
            st.caption(
                "Collect snippets from candidates or jot down notes to reuse later."
            )
            st.text_area(
                "Snippet clipboard",
                key="textual_refiner_snippet",
                height=140,
            )

            st.subheader("Merge workspace")
            if st.session_state.textual_refiner_merge_sources:
                st.caption(
                    "Sources: "
                    + ", ".join(
                        dict.fromkeys(
                            st.session_state.textual_refiner_merge_sources
                        )
                    )
                )
            st.text_area(
                "Merge workspace",
                key="textual_refiner_merge_text",
                height=180,
                help="Add sentences from candidates or type custom edits, then finalize the merge.",
            )

            merge_cols = st.columns(2)
            with merge_cols[0]:
                st.button(
                    "✅ Finalize merge",
                    key="textual_refiner_finalize_merge_button",
                    disabled=not st.session_state.textual_refiner_merge_text.strip(),
                    on_click=_finalize_merge,
                )
            with merge_cols[1]:
                st.button(
                    "🧹 Clear merge",
                    on_click=_clear_merge_workspace,
                )

        if st.session_state.textual_refiner_instruction_output:
            st.markdown("---")
            st.subheader("Instruction proposal")
            st.markdown(
                st.session_state.textual_refiner_instruction_diff,
                unsafe_allow_html=True,
            )
            proposed_text = st.text_area(
                "AI edited draft",
                value=st.session_state.textual_refiner_instruction_output,
                key="textual_refiner_instruction_editor",
                height=220,
            )
            st.session_state.textual_refiner_instruction_output = proposed_text
            action_cols = st.columns(2)
            with action_cols[0]:
                st.button(
                    "✅ Apply edit",
                    type="primary",
                    on_click=_apply_instruction_edit,
                )
            with action_cols[1]:
                st.button(
                    "↩️ Discard proposal",
                    on_click=_discard_instruction_edit,
                )

        st.markdown("---")
        st.subheader("Candidate rewrites")
        if not st.session_state.textual_refiner_candidates:
            st.info(
                "Select a rewrite style and generate suggestions to compare against your draft."
            )
        else:
            for candidate in st.session_state.textual_refiner_candidates:
                preset = _get_preset(candidate["key"])
                with st.container():
                    st.markdown(f"#### {candidate['label']}")
                    if preset:
                        st.caption(preset["instruction"])
                    if candidate.get("focus"):
                        st.caption(f"Focus: {candidate['focus']}")
                    st.markdown(
                        generate_inline_diff_html(
                            candidate["source"], candidate["text"]
                        ),
                        unsafe_allow_html=True,
                    )
                    st.text_area(
                        "Candidate text",
                        value=candidate["text"],
                        key=f"textual_refiner_candidate_text_{candidate['id']}",
                        height=220,
                    )

                    sentences = _split_sentences(candidate["text"])
                    selection_key = (
                        f"textual_refiner_merge_select_{candidate['id']}"
                    )
                    if sentences:
                        st.multiselect(
                            "Select sentences to add to the merge workspace",
                            options=sentences,
                            key=selection_key,
                        )

                    action_cols = st.columns(3)
                    with action_cols[0]:
                        st.button(
                            "✅ Accept rewrite",
                            key=f"textual_refiner_accept_{candidate['id']}",
                            on_click=_accept_candidate,
                            args=(candidate["id"],),
                        )
                    with action_cols[1]:
                        if st.button(
                            "📋 To clipboard",
                            key=f"textual_refiner_snippet_{candidate['id']}",
                        ):
                            st.session_state.textual_refiner_snippet = candidate[
                                "text"
                            ]
                            st.session_state.textual_refiner_feedback = (
                                "Rewrite copied to the snippet clipboard."
                            )
                            st.rerun()
                    with action_cols[2]:
                        if sentences and st.button(
                            "➕ Add to merge",
                            key=f"textual_refiner_merge_{candidate['id']}",
                        ):
                            _add_sentences_to_merge(candidate["id"], selection_key)
                            st.session_state.textual_refiner_feedback = (
                                "Selected sentences added to the merge workspace."
                            )
                            st.rerun()

    with tab_diagnostics:
        if not diagnostics_enabled:
            st.info(
                "Enable the diagnostics panel from Settings to view rewrite metrics and exports."
            )
        else:
            st.subheader("Quality metrics")
            current_text = st.session_state.textual_refiner_draft
            current_metrics = calculate_text_metrics(current_text)

            last_metrics_dict = st.session_state.get(
                "textual_refiner_last_metrics"
            )
            prev_metrics_dict = st.session_state.get(
                "textual_refiner_prev_metrics"
            )
            last_text = st.session_state.get("textual_refiner_last_text")

            previous_dict: Optional[Dict[str, float]] = None
            if last_metrics_dict is None:
                previous_dict = None
            elif current_text != last_text:
                previous_dict = last_metrics_dict
            else:
                previous_dict = prev_metrics_dict

            previous_metrics = (
                TextMetrics.from_dict(previous_dict) if previous_dict else None
            )

            metrics_entries = _build_metric_entries(current_metrics, previous_metrics)

            metric_columns = st.columns(2)
            for index, entry in enumerate(metrics_entries):
                with metric_columns[index % 2]:
                    st.metric(entry["label"], entry["value"], entry.get("delta"))

            st.caption(
                "Flesch Reading Ease scores range from 0 (difficult) to 100 (very easy)."
            )

            if last_text != current_text:
                st.session_state.textual_refiner_prev_metrics = (
                    last_metrics_dict or current_metrics.to_dict()
                )
                st.session_state.textual_refiner_last_text = current_text
            st.session_state.textual_refiner_last_metrics = current_metrics.to_dict()

            st.markdown("---")
            st.subheader("Quality check")
            if not quality_checks_enabled:
                st.info(
                    "Enable AI quality checks from Settings to run grammar diagnostics on demand."
                )
            else:
                check_disabled = not current_text.strip()
                if st.button(
                    "🔍 Run quality check",
                    key="textual_refiner_run_check",
                    disabled=check_disabled,
                ):
                    with st.spinner("Reviewing draft against style guides…"):
                        try:
                            messages = _build_check_messages(current_text)
                            stream = openai_api_call(messages, "Normal")  # type: ignore[arg-type]
                            st.session_state.textual_refiner_quality_checks = (
                                "".join(list(stream)).strip()
                            )
                        except Exception as exc:  # pylint: disable=broad-except
                            st.error(f"Quality check failed: {exc}")

                if st.session_state.textual_refiner_quality_checks:
                    st.markdown(st.session_state.textual_refiner_quality_checks)
                else:
                    st.info(
                        "Run a quality check to generate an AI-powered diagnostic report."
                    )

            st.markdown("---")
            st.subheader("Export")
            st.checkbox(
                "Include citations for applied fixes",
                key="textual_refiner_export_include_citations",
                help="When enabled, applied fixes are labeled as [Fix n] in the export.",
            )

            style_guides = st.session_state.get("active_style_guides", [])
            include_citations = (
                st.session_state.textual_refiner_export_include_citations
            )
            history_entries = st.session_state.get(
                "textual_refiner_history", []
            )

            docx_bytes = _create_docx_report(
                text=current_text,
                metrics_entries=metrics_entries,
                quality_check=st.session_state.textual_refiner_quality_checks,
                style_guides=style_guides,
                include_citations=include_citations,
                history=history_entries,
            )
            markdown_report = _create_markdown_report(
                text=current_text,
                metrics_entries=metrics_entries,
                quality_check=st.session_state.textual_refiner_quality_checks,
                style_guides=style_guides,
                include_citations=include_citations,
                history=history_entries,
            )

            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            st.download_button(
                "⬇️ Download report (.docx)",
                data=docx_bytes,
                file_name=f"arcana_textual_refiner_report_{timestamp}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            st.download_button(
                "⬇️ Download summary (.md)",
                data=markdown_report,
                file_name=f"arcana_textual_refiner_report_{timestamp}.md",
                mime="text/markdown",
            )

    with tab_history:
        st.subheader("Revision snapshots")
        versions = st.session_state.textual_refiner_versions
        if not versions:
            st.info(
                "Snapshots are created whenever you apply a rewrite or instruction-driven edit."
            )
        else:
            st.write(f"**Stored versions:** {len(versions)}")
            for reverse_index, version in enumerate(reversed(versions)):
                version_index = len(versions) - 1 - reverse_index
                title = f"{version['timestamp']} — {version['label']}"
                with st.expander(title, expanded=False):
                    st.markdown("**Diff vs. current draft**")
                    st.markdown(
                        generate_version_diff_html(
                            st.session_state.textual_refiner_draft,
                            version["text"],
                        ),
                        unsafe_allow_html=True,
                    )
                    st.markdown("**Snapshot text**")
                    st.write(version["text"])
                    col_revert, col_copy = st.columns([1, 1])
                    with col_revert:
                        st.button(
                            "⏪ Revert to this version",
                            key=f"textual_refiner_revert_{version_index}",
                            on_click=_queue_revert,
                            args=(version_index,),
                        )
                    with col_copy:
                        if st.button(
                            "📋 Copy to clipboard",
                            key=f"textual_refiner_copy_version_{version_index}",
                        ):
                            st.session_state.textual_refiner_snippet = version["text"]
                            st.session_state.textual_refiner_feedback = (
                                "Snapshot copied to the snippet clipboard."
                            )
                            st.rerun()

        st.markdown("---")
        st.subheader("Applied rewrites")
        if not st.session_state.textual_refiner_history:
            st.info(
                "Accept a rewrite, finalize a merge, or apply an instruction edit to populate this log."
            )
        else:
            for entry in st.session_state.textual_refiner_history:
                title = f"{entry['timestamp']} — {entry['label']}"
                with st.expander(title, expanded=False):
                    if entry.get("focus"):
                        st.caption(f"Focus: {entry['focus']}")
                    st.markdown("**Diff vs. original**")
                    st.markdown(
                        generate_inline_diff_html(
                            entry["source_text"], entry["rewritten_text"]
                        ),
                        unsafe_allow_html=True,
                    )
                    st.markdown("**Rewritten text**")
                    st.write(entry["rewritten_text"])
                    st.markdown("**Original text**")
                    st.write(entry["source_text"])
